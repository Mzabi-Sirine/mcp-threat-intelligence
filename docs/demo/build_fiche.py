#!/usr/bin/env python3
"""Fiche de conduite de la démonstration live (2 pages A4)."""

from __future__ import annotations

import subprocess
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

HERE = Path(__file__).resolve().parent
PROJECT = HERE.parents[1]
OUT = HERE / "Fiche_demo_soutenance_MCP_TI.docx"
LOGO = PROJECT / "figures" / "logoTek-up_university.png"

NAVY = RGBColor(0x10, 0x24, 0x3F)
NAVY_HEX = "10243F"
GOLD = RGBColor(0x8A, 0x6D, 0x1F)
GOLD_HEX = "C9A227"
INK = RGBColor(0x1F, 0x2A, 0x3A)
GREY = RGBColor(0x5A, 0x66, 0x76)
GREEN = RGBColor(0x2E, 0x7D, 0x5B)
RED = RGBColor(0xA6, 0x3B, 0x3B)
PALE = "F3F5F8"
PALE_GOLD = "FBF6E3"
PALE_GREEN = "EDF6F1"
PALE_RED = "FBF0F0"


def font(run, name="Calibri", size=None, bold=None, italic=None, color=None):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs"):
        rfonts.set(qn(attr), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if italic is not None:
        run.font.italic = italic
    if color is not None:
        run.font.color.rgb = color


def shade(cell, fill):
    el = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    el.append(shd)


def borders(par, *, bottom=None, top=None, color="D5DBE3", size=6):
    ppr = par._p.get_or_add_pPr()
    pbdr = ppr.find(qn("w:pBdr"))
    if pbdr is None:
        pbdr = OxmlElement("w:pBdr")
        ppr.append(pbdr)
    for tag, use in (("top", top), ("bottom", bottom)):
        if not use:
            continue
        el = OxmlElement(f"w:{tag}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(size))
        el.set(qn("w:space"), "3")
        el.set(qn("w:color"), color)
        pbdr.append(el)


def space(par, before=0, after=0, line=1.0):
    pf = par.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line


def fixed_table(t, widths):
    t.autofit = False
    for row in t.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = width
    grid = t._tbl.find(qn("w:tblGrid"))
    if grid is not None:
        for col, width in zip(grid.findall(qn("w:gridCol")), widths):
            col.set(qn("w:w"), str(int(width.twips)))


def make_table(doc, headers, rows, widths, font_size=8.6, header_fill=NAVY_HEX):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.style = "Table Grid"
    for index, label in enumerate(headers):
        cell = t.rows[0].cells[index]
        shade(cell, header_fill)
        cell.text = ""
        p = cell.paragraphs[0]
        space(p, before=2, after=2)
        r = p.add_run(label)
        font(r, size=font_size, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
    for r_index, row in enumerate(rows):
        cells = t.add_row().cells
        for c_index, value in enumerate(row):
            cells[c_index].text = ""
            p = cells[c_index].paragraphs[0]
            space(p, before=2, after=2, line=1.0)
            r = p.add_run(str(value))
            font(r, size=font_size, color=INK)
            if r_index % 2 == 1:
                shade(cells[c_index], PALE)
    fixed_table(t, widths)
    return t


def h1(doc, text):
    p = doc.add_paragraph()
    space(p, before=10, after=4)
    borders(p, bottom=True, color=GOLD_HEX, size=8)
    r = p.add_run(text)
    font(r, size=11.5, bold=True, color=NAVY)
    return p


def body(doc, text, size=9.4, color=INK, before=0, after=3):
    p = doc.add_paragraph()
    space(p, before=before, after=after, line=1.06)
    r = p.add_run(text)
    font(r, size=size, color=color)
    return p


def bullet(doc, lead, text):
    p = doc.add_paragraph(style="List Bullet")
    space(p, before=1, after=1, line=1.03)
    r = p.add_run(lead)
    font(r, size=9.2, bold=True, color=NAVY)
    r2 = p.add_run(" — " + text)
    font(r2, size=9.2, color=INK)


def callout(doc, title, text, fill=PALE_GOLD, accent=GOLD):
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = t.rows[0].cells[0]
    shade(cell, fill)
    cell.text = ""
    p = cell.paragraphs[0]
    space(p, before=4, after=1)
    r = p.add_run(title)
    font(r, size=9.0, bold=True, color=accent)
    p2 = cell.add_paragraph()
    space(p2, before=0, after=4, line=1.04)
    r2 = p2.add_run(text)
    font(r2, size=8.9, color=INK)
    fixed_table(t, [Cm(17.0)])


def main() -> int:
    doc = Document()
    section = doc.sections[0]
    section.page_height = Cm(29.7)
    section.page_width = Cm(21.0)
    section.top_margin = Cm(1.5)
    section.bottom_margin = Cm(1.4)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)

    header = section.header.paragraphs[0]
    header.text = ""
    runs = [
        ("TEK-UP SSIR 2026-2027", True, NAVY),
        ("   ·   Démonstration live — Serveur MCP de Threat Intelligence", False, GREY),
    ]
    for text, bold, color in runs:
        r = header.add_run(text)
        font(r, size=8.0, bold=bold, color=color)
    borders(header, bottom=True, color="D5DBE3", size=4)

    p = doc.add_paragraph()
    space(p, before=4, after=1)
    r = p.add_run("Démonstration devant le professeur")
    font(r, size=17, bold=True, color=NAVY)
    p2 = doc.add_paragraph()
    space(p2, before=0, after=2)
    r2 = p2.add_run("Serveur MCP intelligent de Threat Intelligence · 7 scènes · environ 6 minutes")
    font(r2, size=10.5, color=GOLD, bold=True)
    p3 = doc.add_paragraph()
    space(p3, before=0, after=6)
    r3 = p3.add_run("Mzabi Sirine · Khaled Ellif   |   encadrant : Tarek Hdji   |   données synthétiques, laboratoire local")
    font(r3, size=8.6, color=GREY)

    callout(doc, "À faire 2 minutes avant",
            "Ouvrir un terminal WSL, se placer dans le projet du Bureau puis lancer :  "
            "cd /mnt/c/Users/dell/Desktop/art/projet  &&  bash demo/lancer_demo.sh.  "
            "Variante atelier : cd /root/tekup-mcp-ti && python3 scripts/demo_defense.py --present.  "
            "La démo utilise un laboratoire dédié (preuves/demo_prof) : elle ne modifie pas les preuves du rapport. "
            "Appuyer sur Entrée après chaque scène.", fill=PALE_GREEN, accent=GREEN)

    h1(doc, "1. Ce que vous dites, scène par scène")
    rows = [
        ("S1 · Découverte",
         "Protocole 2025-06-18, serveur v1.0.0, six outils tous en lecture seule, volume de bulletins visibles.",
         "« Le serveur expose uniquement des outils de lecture ; aucune écriture n'est possible depuis MCP. »", "40 s"),
        ("S2 · Fuite bloquée",
         "Jeton Confidentiel sur un bulletin Secret : refus, canari absent, journal deny / HABILITATION_INSUFFISANTE.",
         "« Le niveau d'habilitation ne suffit pas : la demande est refusée et le canari ne fuit pas. »", "50 s"),
        ("S3 · Cloisonnement",
         "Jeton Secret sans CRYPTO : refus, raison BESOIN_CONNAITRE_NON_SATISFAIT.",
         "« Même Secret, il manque le compartiment : le besoin d'en connaître est vérifié séparément. »", "45 s"),
        ("S4 · Accès autorisé",
         "Jeton Secret + CRYPTO : bulletin livré, canari présent, décision allow.",
         "« Le contrôle positif fonctionne : le canari apparaît seulement avec les bons droits. »", "40 s"),
        ("S5 · Injection",
         "Bulletin ADV-2026-012 : fanions appel_outil, balise_active, exfiltration, ignorer_instructions ; instruction brute absente, marqueur de neutralisation présent.",
         "« L'instruction hostile est signalée et neutralisée ; elle n'est jamais exécutée ni livrée telle quelle. »", "50 s"),
        ("S6 · Audit chaîné",
         "Journal valide, puis copie falsifiée : valide=False, rupture localisée, raison « empreinte invalide ».",
         "« Le journal est chaîné : toute modification est détectée à la séquence près. »", "50 s"),
        ("S7 · Avant / après",
         "Maquette naïve : canari=True et source en clair=True ; serveur durci : False / False ; contrôle positif True.",
         "« Mêmes bulletins, deux serveurs : la maquette fuit, le serveur durci masque, refuse et journalise. »", "60 s"),
    ]
    make_table(doc, ["Scène", "Ce qui s'affiche", "Ce que vous dites", "Durée"], rows,
               [Cm(2.6), Cm(6.7), Cm(6.3), Cm(1.4)], font_size=7.9)

    h1(doc, "2. Trois phrases à garder")
    bullet(doc, "Preuve, pas promesse",
           "le canari, le refus et le journal rendent le contrôle vérifiable à l'écran.")
    bullet(doc, "Deux verrous",
           "l'habilitation donne le niveau, le compartiment donne le besoin d'en connaître ; les deux sont contrôlés.")
    bullet(doc, "Journal d'abord",
           "chaque appel — accepté ou refusé — laisse une entrée chaînée ; une falsification casse la chaîne.")

    h1(doc, "3. Questions probables du jury")
    qa = [
        ("Pourquoi MCP plutôt qu'une API REST ?",
         "MCP standardise la découverte et l'appel d'outils en JSON-RPC 2.0 ; les annotations readOnlyHint rendent le "
         "caractère lecture seule vérifiable par le client. Ici le transport est stdio local ; HTTP et mTLS sont des "
         "limites assumées, pas des acquis."),
        ("Comment prouvez-vous qu'un bulletin Secret ne fuit pas ?",
         "Chaque bulletin Secret porte un canari unique. Il est absent des réponses refusées (S2, S3) et présent "
         "uniquement avec Secret + CRYPTO (S4). La décision et sa raison sont dans le journal chaîné."),
        ("Quelles sont les limites ?",
         "Transport stdio local, pas de HTTP/mTLS, SQLite mono-nœud, jeu synthétique de 12 bulletins et 30 indicateurs, "
         "pas d'EDR réel. Le rapport nomme huit limites, chacune avec sa condition de levée."),
    ]
    for question, answer in qa:
        p = doc.add_paragraph()
        space(p, before=3, after=0)
        r = p.add_run("Q. " + question)
        font(r, size=9.0, bold=True, color=NAVY)
        p2 = doc.add_paragraph()
        space(p2, before=0, after=2, line=1.03)
        r2 = p2.add_run(answer)
        font(r2, size=8.8, color=INK)

    h1(doc, "4. Clôture et plan de secours")
    body(doc, "Phrase de clôture : « Le serveur ne promet pas la sécurité : il la rend vérifiable — chaque refus et "
              "chaque accès laissent une trace chaînée. »", size=9.2)
    callout(doc, "Si la démonstration ne démarre pas",
            "Ouvrir le rapport, section 6.4 pages 24 à 27 (sept captures réelles), puis les PNG du dossier "
            "projet/captures. La vidéo Demo_MCP_Threat_Intelligence.mp4 (1 min 55, voix off) rejoue les sept scènes ; "
            "Demo_Architecture_MCP_TI.mp4 (1 min 37) présente les quatre zones et les huit points de contrôle ; "
            "Demo_Tests_MCP_TI.mp4 (34 s) filme la campagne réelle : 31 scénarios, puis vérification de la chaîne d'audit. "
            "En direct sur le poste : double-clic Lancer_Tests_Windows.bat (synthèse + audit, fenêtre maintenue). "
            "Notes simples sous chaque diapo : Deck_Soutenance_MCP_TI_notes.pdf.",
            fill=PALE_RED, accent=RED)

    doc.save(OUT)
    print(f"[docx] {OUT}")
    subprocess.run(["libreoffice", "--headless", "--convert-to", "pdf", "--outdir", str(HERE), str(OUT)],
                   check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    print(f"[pdf ] {OUT.with_suffix('.pdf')}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
