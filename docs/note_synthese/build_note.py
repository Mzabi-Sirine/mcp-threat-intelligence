#!/usr/bin/env python3
"""Note de synthèse indépendante : ce qui est testé, ce qui ne l'est pas, limites."""

from __future__ import annotations

import subprocess
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

HERE = Path(__file__).resolve().parent
OUT = HERE / "Note_de_synthese_MCP_TI.docx"

NAVY = RGBColor(0x10, 0x24, 0x3F)
NAVY_HEX = "10243F"
GOLD = RGBColor(0x8A, 0x6D, 0x1F)
GOLD_HEX = "C9A227"
INK = RGBColor(0x1F, 0x2A, 0x3A)
GREY = RGBColor(0x5A, 0x66, 0x76)
GREEN = RGBColor(0x2E, 0x7D, 0x5B)
RED = RGBColor(0xA6, 0x3B, 0x3B)
BLUE = RGBColor(0x2B, 0x5C, 0x8A)
PALE = "F3F5F8"
PALE_GOLD = "FBF6E3"
PALE_GREEN = "EDF6F1"
PALE_RED = "FBF0F0"


def font(run, name="Calibri", size=None, bold=None, italic=None, color=None):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs"):
        rfonts.set(qn(attr), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if italic is not None:
        run.font.italic = italic
    if color is not None:
        run.font.color.rgb = color


def shade(cell, fill):
    el = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    el.append(shd)


def borders(par, *, bottom=None, top=None, color="D5DBE3", size=6):
    ppr = par._p.get_or_add_pPr()
    pbdr = ppr.find(qn("w:pBdr"))
    if pbdr is None:
        pbdr = OxmlElement("w:pBdr")
        ppr.append(pbdr)
    for tag, use in (("top", top), ("bottom", bottom)):
        if not use:
            continue
        el = OxmlElement(f"w:{tag}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(size))
        el.set(qn("w:space"), "3")
        el.set(qn("w:color"), color)
        pbdr.append(el)


def space(par, before=0, after=0, line=1.0):
    pf = par.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line


def field(par, instr, cached="1"):
    r = par.add_run()
    fld = OxmlElement("w:fldChar")
    fld.set(qn("w:fldCharType"), "begin")
    r._r.append(fld)
    r2 = par.add_run()
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = f" {instr} "
    r2._r.append(it)
    r3 = par.add_run()
    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    r3._r.append(sep)
    r4 = par.add_run(cached)
    font(r4, size=7.8, color=GREY)
    r5 = par.add_run()
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    r5._r.append(end)


def fixed_table(t, widths):
    t.autofit = False
    for row in t.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = width
    grid = t._tbl.find(qn("w:tblGrid"))
    if grid is not None:
        for col, width in zip(grid.findall(qn("w:gridCol")), widths):
            col.set(qn("w:w"), str(int(width.twips)))


def no_split(row):
    trpr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:cantSplit")
    trpr.append(el)


def rich(par, text, size=9.6, color=INK, bold=False, italic=False):
    r = par.add_run(text)
    font(r, size=size, color=color, bold=bold, italic=italic)


def h1(doc, text):
    p = doc.add_paragraph()
    space(p, before=10, after=4)
    r = p.add_run(text)
    font(r, size=14.5, bold=True, color=NAVY)
    borders(p, bottom=True, color=GOLD_HEX, size=8)
    return p


def h2(doc, text):
    p = doc.add_paragraph()
    space(p, before=8, after=3)
    r = p.add_run(text)
    font(r, size=10.8, bold=True, color=NAVY)
    return p


def body(doc, text, size=9.8, after=4, color=INK, italic=False):
    p = doc.add_paragraph()
    space(p, before=0, after=after, line=1.04)
    rich(p, text, size=size, color=color, italic=italic)
    return p


def bullet(doc, label, text, size=9.6):
    p = doc.add_paragraph()
    space(p, before=1, after=1, line=1.04)
    p.paragraph_format.left_indent = Cm(0.45)
    p.paragraph_format.first_line_indent = Cm(-0.45)
    r = p.add_run("• ")
    font(r, size=size, color=GOLD, bold=True)
    if label:
        r = p.add_run(label + " — ")
        font(r, size=size, bold=True, color=NAVY)
    rich(p, text, size=size)


def make_table(doc, headers, rows, widths, *, font_size=8.7, header_fill=NAVY_HEX,
               first_col_bold=True, band=True, row_colors=None, center_cols=()):
    t = doc.add_table(rows=len(rows) + 1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for j, head in enumerate(headers):
        cell = t.rows[0].cells[j]
        cell.text = ""
        shade(cell, header_fill)
        p = cell.paragraphs[0]
        space(p, before=2, after=2)
        r = p.add_run(head)
        font(r, size=font_size - 0.2, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
        if j in center_cols:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for i, row in enumerate(rows, 1):
        for j, val in enumerate(row):
            cell = t.rows[i].cells[j]
            cell.text = ""
            fill = None
            if row_colors and row_colors[i - 1]:
                fill = row_colors[i - 1][j]
            elif band and i % 2 == 0:
                fill = PALE
            if fill:
                shade(cell, fill)
            p = cell.paragraphs[0]
            space(p, before=2, after=2, line=1.0)
            if j in center_cols:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(str(val))
            font(r, size=font_size, bold=(first_col_bold and j == 0),
                 color=NAVY if j == 0 and first_col_bold else INK)
        no_split(t.rows[i])
    fixed_table(t, widths)
    t.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    return t


def page_setup(doc):
    sec = doc.sections[0]
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    sec.top_margin = Cm(1.5)
    sec.bottom_margin = Cm(1.4)
    sec.left_margin = Cm(1.8)
    sec.right_margin = Cm(1.8)
    sec.header_distance = Cm(0.7)
    sec.footer_distance = Cm(0.6)


def header_footer(doc):
    sec = doc.sections[0]
    sec.different_first_page_header_footer = True
    hp = sec.header.paragraphs[0]
    hp.text = ""
    borders(hp, bottom=True, color="D5DBE3", size=4)
    hp.paragraph_format.tab_stops.clear_all()
    hp.paragraph_format.tab_stops.add_tab_stop(Cm(17.4), WD_TAB_ALIGNMENT.RIGHT)
    r = hp.add_run("Note de synthèse — MCP Threat Intelligence")
    font(r, size=7.8, color=GREY)
    r = hp.add_run("\tTEK-UP · SSIR · 2026")
    font(r, size=7.8, color=GREY)
    fp = sec.footer.paragraphs[0]
    fp.text = ""
    fp.paragraph_format.tab_stops.clear_all()
    fp.paragraph_format.tab_stops.add_tab_stop(Cm(17.4), WD_TAB_ALIGNMENT.RIGHT)
    r = fp.add_run("Mzabi Sirine · Khaled Ellif — dossier preuves/ joint")
    font(r, size=7.5, color=GREY)
    r = fp.add_run("\t")
    font(r, size=7.5, color=GREY)
    field(fp, "PAGE", "1")
    r = fp.add_run(" / ")
    font(r, size=7.5, color=GREY)
    field(fp, "NUMPAGES", "4")


def title_band(doc):
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    fixed_table(t, [Cm(17.4)])
    cell = t.rows[0].cells[0]
    shade(cell, NAVY_HEX)
    cell.text = ""
    p = cell.paragraphs[0]
    space(p, before=12, after=2)
    r = p.add_run("NOTE DE SYNTHÈSE  ·  VERSION 1.0")
    font(r, size=8.4, bold=True, color=RGBColor(0xE8, 0xC9, 0x6A))
    p = cell.add_paragraph()
    space(p, before=2, after=2)
    r = p.add_run("Ce qui est testé,\nce qui ne l'est pas")
    font(r, size=25, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
    p.paragraph_format.line_spacing = 1.02
    p = cell.add_paragraph()
    space(p, before=4, after=12)
    r = p.add_run("Serveur MCP intelligent de renseignement sur les menaces — TEK-UP SSIR")
    font(r, size=10.5, color=RGBColor(0xC9, 0xD5, 0xE3))


def kpis(doc):
    cards = [
        ("31 / 31", "scénarios réussis", "0 échec", GREEN),
        ("16 / 16", "abus maîtrisés", "100 % de détection, aucune fuite", GREEN),
        ("12,42 ms", "latence p95", "60 requêtes · cible ≤ 50 ms", NAVY),
        ("8 → 0", "faiblesses comparées", "maquette naïve → serveur durci", GOLD),
    ]
    t = doc.add_table(rows=1, cols=4)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    fixed_table(t, [Cm(4.35)] * 4)
    for cell, (big, label, detail, accent) in zip(t.rows[0].cells, cards):
        shade(cell, PALE)
        cell.text = ""
        p = cell.paragraphs[0]
        space(p, before=4, after=0)
        r = p.add_run(big)
        font(r, size=19, bold=True, color=accent)
        p = cell.add_paragraph()
        space(p, before=0, after=1)
        r = p.add_run(label)
        font(r, size=8.8, bold=True, color=NAVY)
        p = cell.add_paragraph()
        space(p, before=0, after=4)
        r = p.add_run(detail)
        font(r, size=7.4, color=GREY)


def build():
    doc = Document()
    page_setup(doc)
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(9.8)
    style.paragraph_format.space_after = Pt(0)
    header_footer(doc)
    settings = doc.settings.element
    upd = OxmlElement("w:updateFields")
    upd.set(qn("w:val"), "true")
    settings.append(upd)

    title_band(doc)

    p = doc.add_paragraph()
    space(p, before=8, after=6)
    rich(p, "Cette note ne remplace pas le rapport : elle isole ce qui a été réellement exécuté, "
            "ce qui reste déclaratif et ce qui n'a pas été testé. Elle est écrite pour le jury, "
            "puis pour l'équipe qui reprendra le projet : chaque affirmation renvoie à un fichier "
            "de preuve ou à une commande rejouable.", size=9.9)

    kpis(doc)

    h1(doc, "1. Résultat d'ensemble")
    body(doc, "La campagne finale du 27 octobre 2026 a exécuté 31 scénarios sur le serveur durci. "
              "Les chiffres ci-dessous sont ceux du journal de campagne, pas des estimations.")

    rows = [
        ("Scénarios", "31 / 31", "14 nominaux · 1 durcissement · 16 abus", "preuves/logs/tests_latest.json"),
        ("Abus", "16 / 16", "100 % refusés ou neutralisés", "mêmes preuves, catégorie abus"),
        ("Latence", "p95 12,42 ms", "n = 60, audit actif, cible ≤ 50 ms", "métriques de campagne"),
        ("Révocation", "10,0 ms", "sur la requête suivante", "commande mesurée à 121,31 ms"),
        ("Audit", "115 / 115", "116 entrées, chaîne valide", "une erreur de protocole incluse"),
        ("Injection", "3 / 3", "0 faux positif sur 11 bulletins bénins", "corpus unitaire"),
        ("Comparaison", "8 → 0", "mêmes 8 contrôles, maquette puis serveur", "baseline_naive.py"),
        ("Durcissement", "uid 65534", "base 0440, clé 0400, 0 port", "setpriv, aucune écoute"),
        ("Données", "12 · 30 · 6", "bulletins, indicateurs, actifs synthétiques", "init-db, aucune donnée réelle"),
        ("Outils", "6", "lecture seule, découverte MCP", "tools/list, readOnlyHint"),
    ]
    make_table(doc, ["Indicateur", "Valeur", "Lecture", "Source"], rows,
               [Cm(3.0), Cm(2.8), Cm(6.4), Cm(5.2)], font_size=8.4, center_cols=(1,))

    p = doc.add_paragraph()
    space(p, before=6, after=0)
    rich(p, "Lecture : le projet ne se contente pas de « fonctionner » ; il échoue proprement. "
            "Les 16 abus sont refusés avec un message générique, journalisés avec un motif distinct, "
            "et le journal reste vérifiable.", size=9.1, italic=True, color=GREY)

    h1(doc, "2. Statuts de livraison")
    p = doc.add_paragraph()
    space(p, before=2, after=3)
    r = p.add_run("Statuts utilisés dans ce document : ")
    font(r, size=8.4, bold=True, color=GREY)
    for label, color in (("testé", GREEN), ("configuré", BLUE), ("proposé", GOLD), ("déclaratif", RED)):
        r = p.add_run(f"■ {label}   ")
        font(r, size=8.4, bold=True, color=color)

    rows = [
        ("Testé", "serveur, ABAC, masquage, audit, révocation, durcissement", "31/31, 115/115, 8 → 0"),
        ("Configuré", "base synthétique, 6 outils, journal local, sauvegardes", "12/30/6, init-db, backups"),
        ("Proposé", "systemd, isolation réseau, HTTP/OAuth 2.1/mTLS, ancrage externe", "écrits, non exécutés en production"),
        ("Déclaratif", "agent EDR, chiffrement de disque, posture", "signal déclaré, non attesté"),
    ]
    make_table(doc, ["Statut", "Périmètre", "Preuve ou réserve"], rows,
               [Cm(2.4), Cm(8.9), Cm(6.1)], font_size=8.6,
               row_colors=[[c] * 3 for c in (PALE_GREEN, PALE, PALE_GOLD, PALE_RED)])

    p = doc.add_paragraph()
    space(p, before=5, after=0)
    rich(p, "Un statut n'est ni une intention ni une promesse : « testé » signifie rejoué et mesuré ; "
            "« proposé » signifie écrit mais pas exécuté dans ce laboratoire.", size=9.1, italic=True, color=GREY)

    doc.add_page_break()

    h1(doc, "3. Ce qui est testé")
    body(doc, "Les 31 scénarios se répartissent en huit familles. Chaque famille cite les exigences "
              "qu'elle valide ; le détail ligne à ligne figure dans le rapport et dans le JSON de campagne.")

    rows = [
        ("Protocole et outils", "6", "initialize, 6 outils en lecture seule, ressources filtrées, entrées invalides", "EF-01, EF-10, EF-12"),
        ("Habilitation et cloisonnement", "10", "niveaux 0 à 3, compartiments, IDOR, rôle et finalité, recherche", "EF-03, EF-06, EF-11"),
        ("Masquage et injection", "2", "source Secret masquée, canari, consigne hostile neutralisée", "EF-04, EF-05, ENF-12"),
        ("Renseignement et score", "3", "décroissance d'un indicateur, corrélation de famille, score explicable", "EF-06, EF-07"),
        ("Jetons et posture", "5", "jeton absent, expiré, altéré, révoqué ; poste non conforme", "EF-02, EF-09, ENF-04"),
        ("Audit et persistance", "2", "redémarrage sans perte, falsification détectée", "EF-08, ENF-08"),
        ("Réseau, débit, dépendances", "2", "aucun client réseau, limiteur de débit, stdlib uniquement", "EF-13, ENF-05, ENF-10"),
        ("Durcissement", "1", "uid 65534, base 0440, clé 0400, écriture refusée", "ENF-06, ENF-07"),
    ]
    make_table(doc, ["Famille", "N", "Ce qui est prouvé", "Exigences"], rows,
               [Cm(4.0), Cm(1.0), Cm(8.0), Cm(4.4)], font_size=8.6, center_cols=(1,))

    h2(doc, "Couverture des exigences")
    body(doc, "24 des 25 exigences (13 fonctionnelles, 12 non fonctionnelles) sont citées par au moins "
              "un scénario. La 25e, ENF-03 — 100 % des appels attendus journalisés — est vérifiée par "
              "la mesure d'audit et non par un scénario : 115 appels sur 115, 116 entrées dont une "
              "erreur de protocole, chaîne valide. Aucune exigence n'est laissée sans preuve.")
    body(doc, "Trois scénarios méritent d'être cités au jury : A06 et A07, qui tentent deux fuites "
              "distinctes — habilitation insuffisante puis défaut de cloisonnement — et A14, qui "
              "falsifie une copie du journal et vérifie que la rupture de séquence est détectée.",
         size=9.4)

    h2(doc, "Ce que le jury peut vérifier en séance")
    bullet(doc, "Fuite", "le jeton Confidentiel lit un bulletin Secret : refus générique, deux motifs distincts dans le journal.")
    bullet(doc, "Injection", "un bulletin partenaire contient une consigne hostile : elle est signalée et neutralisée avant livraison.")
    bullet(doc, "Audit", "la commande verify-audit renvoie ok=true, entries=116 ; une copie falsifiée casse la chaîne.")
    bullet(doc, "Durcissement", "le processus tourne sous uid 65534, la base est en 0440, la clé en 0400, aucun port n'écoute.")

    h2(doc, "Repères horodatés de la campagne")
    rows = [
        ("13:08", "Premier essai : échec réel de l'audit sur la deuxième écriture", "iteration_01_echec_audit.txt"),
        ("13:14", "Correction du décodage dans audit._last, puis rejeu", "même trace"),
        ("13:16", "Campagne complète : 31/31, 16/16 abus", "tests_latest.json"),
        ("13:21", "Démonstration isolée : fuite bloquée", "demo_transcript.txt"),
        ("14:21", "Captures et figures régénérées", "preuves/captures, figures"),
    ]
    make_table(doc, ["Heure", "Étape", "Preuve"], rows,
               [Cm(1.5), Cm(9.5), Cm(6.4)], font_size=8.4)

    doc.add_page_break()

    h1(doc, "4. Ce qui n'est pas testé")
    body(doc, "Une limite nommée n'est pas un échec caché : c'est une condition de production identifiée. "
              "Le tableau donne la conséquence réelle et la façon de lever chaque limite.")

    rows = [
        ("Isolation réseau par espace de noms", "proposé",
         "le test n'a pas été exécuté ; l'absence de client réseau est vérifiée statiquement",
         "lancer le serveur sans interface et vérifier zéro socket"),
        ("Agent EDR et chiffrement de disque", "déclaratif",
         "la posture est un signal déclaré, pas une preuve cryptographique",
         "agent signé, attestation vérifiée, test de challenge"),
        ("Ancrage externe du journal", "absent",
         "une réécriture complète par un attaquant root n'est pas exclue",
         "horodatage signé externe ou support WORM"),
        ("HTTP, OAuth 2.1, mTLS", "proposé",
         "le transport est stdio local, sans port ni fédération d'identité",
         "passerelle HTTP, mTLS, OAuth, tests de rejeu"),
        ("Charge et concurrence", "non testé",
         "60 requêtes locales sur un seul poste, pas de multi-clients",
         "banc de charge, SQLite WAL ou serveur dédié"),
        ("Données réelles", "hors périmètre",
         "jeu synthétique RFC 5737 et .example, aucune donnée client",
         "anonymisation et accord de traitement"),
        ("Détection d'injection", "heuristique",
         "3/3 sur le corpus, 0 faux positif sur 11, mais pas de garantie sémantique",
         "corpus adversarial élargi, classification dédiée"),
        ("Unité systemd", "proposé",
         "le service tourne par setpriv en laboratoire ; l'unité n'est pas installée",
         "installer, activer, redémarrer, vérifier l'uid"),
    ]
    colors = []
    for _, status, _, _ in rows:
        colors.append([None, None, None, None])
    make_table(doc, ["Limite", "Statut", "Conséquence", "Condition de levée"], rows,
               [Cm(3.7), Cm(2.4), Cm(6.5), Cm(4.8)], font_size=8.3, center_cols=(1,))

    h2(doc, "Priorités avant toute production")
    bullet(doc, "Intégrité", "ancrer le journal à l'extérieur du poste, par horodatage signé ou support WORM.")
    bullet(doc, "Transport", "ajouter une passerelle HTTP avec mTLS et OAuth 2.1, puis rejouer les 16 abus.")
    bullet(doc, "Confiance", "remplacer la posture déclarative par un agent signé et un test de charge multi-clients.")

    p = doc.add_paragraph()
    space(p, before=6, after=0)
    rich(p, "Ces limites ne s'excusent pas : elles se citent avec leur condition de levée. Le serveur est "
            "sûr dans son périmètre de laboratoire, pas encore en production.", size=9.1, italic=True, color=GREY)

    h2(doc, "Deux questions à anticiper")
    for q, a in [
        ("Le journal est-il inviolable ?",
         "Non. La chaîne détecte une altération locale, pas une réécriture complète par un attaquant root. "
         "C'est pourquoi l'ancrage externe est la première priorité avant production."),
        ("Pourquoi ne pas avoir testé HTTP ?",
         "Le transport stdio local réduit la surface et suffit à la démonstration. HTTP, OAuth 2.1 et mTLS "
         "sont proposés et documentés, pas revendiqués comme testés."),
    ]:
        p = doc.add_paragraph()
        space(p, before=3, after=0)
        r = p.add_run(q)
        font(r, size=9.6, bold=True, color=NAVY)
        p = doc.add_paragraph()
        space(p, before=0, after=2, line=1.03)
        p.paragraph_format.left_indent = Cm(0.35)
        rich(p, a, size=9.3)

    doc.add_page_break()

    h1(doc, "5. Conditions de rejeu")
    body(doc, "Les commandes ci-dessous partent du dossier du projet, sur la même machine que la campagne. "
              "Elles n'utilisent que la bibliothèque standard de Python 3.14.4.")

    code = (
        "cd /root/tekup-mcp-ti\n"
        "python3 tests/run_tests.py\n"
        "# attendu : 31/31, 16/16 abus, p95 ≤ 50 ms\n"
        "python3 scripts/demo_defense.py\n"
        "# attendu : fuite bloquée, Secret autorisé, injection neutralisée\n"
        "python3 src/ti_cli.py verify-audit --path preuves/lab/audit.jsonl\n"
        "# attendu : ok=true, entries=116"
    )
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    fixed_table(t, [Cm(17.4)])
    cell = t.rows[0].cells[0]
    shade(cell, "F4F6F8")
    cell.text = ""
    p = cell.paragraphs[0]
    space(p, before=5, after=1)
    for i, line in enumerate(code.splitlines()):
        if i:
            p = cell.add_paragraph()
            space(p, before=0, after=1)
        r = p.add_run(line)
        font(r, name="Consolas", size=8.8,
             color=GREY if line.startswith("#") else NAVY)

    h2(doc, "Pièces jointes et traces")
    rows = [
        ("tests_latest.json", "31 scénarios, métriques, limites", "preuves/logs/"),
        ("metrics_latest.json", "latence, révocation, audit, injection", "preuves/logs/"),
        ("demo_transcript.txt", "démonstration de fuite bloquée", "preuves/logs/"),
        ("01 à 07_*.png", "7 captures d'écran réelles du terminal", "preuves/captures_ecran/"),
        ("audit.jsonl", "journal chaîné, 116 entrées", "preuves/lab/"),
    ]
    make_table(doc, ["Fichier", "Contenu", "Dossier"], rows,
               [Cm(4.0), Cm(9.0), Cm(4.4)], font_size=8.5)

    h1(doc, "6. Ce que l'équipe retient")
    bullet(doc, "Journal d'abord", "un contrôle non journalisé n'est pas un contrôle : l'échec fermé a dicté la conception du journal avant celle des outils.")
    bullet(doc, "Décider et étiqueter ensemble", "masquer à la sortie ne suffit pas ; la décision d'accès et l'étiquette doivent être calculées au même endroit.")
    bullet(doc, "Garder la maquette", "la version naïve a été conservée pour mesurer les 8 faiblesses, pas seulement pour montrer le résultat final.")

    h1(doc, "7. Dossier remis")
    rows = [
        ("Rapport PFE", "35 pages, DOCX + PDF, sommaire, figures et tableaux numérotés", "docs/rapport/"),
        ("Deck de soutenance", "16 diapositives, PPTX + PDF, notes orales et questions jury", "docs/deck/"),
        ("Conducteur", "19 pages A4, une page par diapositive, plan de secours", "docs/deck/"),
        ("Note de synthèse", "ce document, 4 pages, DOCX + PDF", "docs/note_synthese/"),
        ("Preuves", "7 captures réelles, 31 tests, journaux JSON, audit chaîné", "preuves/"),
    ]
    make_table(doc, ["Livrable", "Contenu", "Dossier"], rows,
               [Cm(3.6), Cm(9.2), Cm(4.6)], font_size=8.4)

    h1(doc, "8. Conclusion")
    body(doc, "Le projet est défendable sur ce qu'il prouve : 31 scénarios, 16 abus maîtrisés, "
              "8 faiblesses neutralisées, 115 appels journalisés sur 115, et un durcissement vérifié. "
              "Il ne prétend pas être prêt pour la production : huit limites sont nommées, dont quatre "
              "directement techniques, et chacune a sa condition de levée. La soutenance doit dire "
              "les deux, sans masquer l'un par l'autre.")

    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    fixed_table(t, [Cm(17.4)])
    cell = t.rows[0].cells[0]
    shade(cell, PALE_GOLD)
    cell.text = ""
    p = cell.paragraphs[0]
    space(p, before=4, after=4)
    r = p.add_run("Phrase à retenir  ")
    font(r, size=8.4, bold=True, color=GOLD)
    r = p.add_run("« Le secret n'est pas une promesse de politique : c'est une propriété testée, "
                  "et les limites restantes sont nommées. »")
    font(r, size=9.8, bold=True, color=NAVY)

    p = doc.add_paragraph()
    space(p, before=16, after=2)
    r = p.add_run("Mzabi Sirine  ·  Khaled Ellif")
    font(r, size=10.5, bold=True, color=NAVY)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    space(p, before=1, after=8)
    r = p.add_run("Encadrant académique : Tarek Hdji")
    font(r, size=9, color=GREY)
    p = doc.add_paragraph()
    space(p, before=0, after=0)
    r = p.add_run("TEK-UP — Spécialité SSIR · 27 octobre 2026 · Version 1.0")
    font(r, size=8.6, color=GREY)

    doc.save(OUT)
    print(f"[docx] {OUT}")
    subprocess.run(
        ["libreoffice", "--headless", "--convert-to", "pdf", "--outdir", str(OUT.parent), str(OUT)],
        check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL,
    )
    pdf = OUT.with_suffix(".pdf")
    print(f"[pdf ] {pdf}")
    info = subprocess.run(["pdfinfo", str(pdf)], capture_output=True, text=True).stdout
    for line in info.splitlines():
        if line.startswith(("Pages:", "Page size:")):
            print("      ", line)


if __name__ == "__main__":
    build()
