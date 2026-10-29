#!/usr/bin/env python3
"""Construit le conducteur de soutenance imprimable (A4) à partir de deck_content.py."""

from __future__ import annotations

import re
import subprocess
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
from deck_content import SLIDES  # noqa: E402

OUT = HERE / "Conducteur_soutenance_MCP_TI.docx"
NAVY = RGBColor(0x10, 0x24, 0x3F)
NAVY_HEX = "10243F"
GOLD = RGBColor(0x8A, 0x6D, 0x1F)
GOLD_HEX = "C9A227"
INK = RGBColor(0x1F, 0x2A, 0x3A)
GREY = RGBColor(0x5A, 0x66, 0x76)
PALE = "F3F5F8"
PALE_GOLD = "FBF6E3"

NUM_RE = re.compile(r"(?<![A-Za-z0-9])(\d+(?:[.,]\d+)?(?:\s?(?:%|ms|s\b|min|cm|points?))?)")


def font(run, name="Calibri", size=None, bold=None, italic=None, color=None):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), name)
    rfonts.set(qn("w:hAnsi"), name)
    rfonts.set(qn("w:cs"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if italic is not None:
        run.font.italic = italic
    if color is not None:
        run.font.color.rgb = color


def shade(cell_or_par, fill):
    if hasattr(cell_or_par, "_tc"):
        el = cell_or_par._tc.get_or_add_tcPr()
    else:
        el = cell_or_par._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    el.append(shd)


def fixed_table(t, widths):
    t.autofit = False
    for row in t.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = width
    grid = t._tbl.find(qn("w:tblGrid"))
    if grid is not None:
        for col, width in zip(grid.findall(qn("w:gridCol")), widths):
            col.set(qn("w:w"), str(int(width.twips)))


def borders(par, *, bottom=None, top=None, left=None, color="D5DBE3", size=6):
    ppr = par._p.get_or_add_pPr()
    pbdr = ppr.find(qn("w:pBdr"))
    if pbdr is None:
        pbdr = OxmlElement("w:pBdr")
        ppr.append(pbdr)
    for tag, use in (("top", top), ("left", left), ("bottom", bottom)):
        if not use:
            continue
        el = OxmlElement(f"w:{tag}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(size))
        el.set(qn("w:space"), "3")
        el.set(qn("w:color"), color)
        pbdr.append(el)


def space(par, before=0, after=0, line=1.0):
    pf = par.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line


def rich(par, text, size=9.4, color=INK, bold_numbers=True, base_bold=False):
    """Écrit le texte en mettant les nombres en gras, comme ils sont prononcés."""
    pos = 0
    for m in NUM_RE.finditer(text):
        if m.start() > pos:
            r = par.add_run(text[pos:m.start()])
            font(r, size=size, bold=base_bold, color=color)
        r = par.add_run(m.group(0))
        font(r, size=size, bold=True, color=NAVY)
        pos = m.end()
    if pos < len(text):
        r = par.add_run(text[pos:])
        font(r, size=size, bold=base_bold, color=color)


def field(par, instr, cached="1"):
    r = par.add_run()
    fld = OxmlElement("w:fldChar")
    fld.set(qn("w:fldCharType"), "begin")
    r._r.append(fld)
    r2 = par.add_run()
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = f" {instr} "
    r2._r.append(it)
    r3 = par.add_run()
    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    r3._r.append(sep)
    r4 = par.add_run(cached)
    font(r4, size=8, color=GREY)
    r5 = par.add_run()
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    r5._r.append(end)


def page_setup(doc):
    sec = doc.sections[0]
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    sec.top_margin = Cm(1.6)
    sec.bottom_margin = Cm(1.5)
    sec.left_margin = Cm(1.8)
    sec.right_margin = Cm(1.8)
    sec.header_distance = Cm(0.8)
    sec.footer_distance = Cm(0.6)


def header_footer(doc, title="Conducteur de soutenance — Serveur MCP de Threat Intelligence"):
    sec = doc.sections[0]
    sec.different_first_page_header_footer = True

    hp = sec.header.paragraphs[0]
    hp.text = ""
    tab = hp.paragraph_format.tab_stops.add_tab_stop(Cm(17.4), WD_ALIGN_PARAGRAPH.RIGHT)
    r = hp.add_run(title)
    font(r, size=7.8, color=GREY)
    r = hp.add_run("\tTEK-UP · SSIR · 2026")
    font(r, size=7.8, color=GREY)
    borders(hp, bottom=True, color="D5DBE3", size=4)

    fp = sec.footer.paragraphs[0]
    fp.text = ""
    fp.paragraph_format.tab_stops.add_tab_stop(Cm(17.4), WD_ALIGN_PARAGRAPH.RIGHT)
    r = fp.add_run("Mzabi Sirine · Khaled Ellif — TEK-UP SSIR")
    font(r, size=7.5, color=GREY)
    r = fp.add_run("\t27 octobre 2026 · ")
    font(r, size=7.5, color=GREY)
    field(fp, "PAGE", "1")
    r = fp.add_run(" / ")
    font(r, size=7.5, color=GREY)
    field(fp, "NUMPAGES", "19")


def cover(doc):
    for _ in range(3):
        p = doc.add_paragraph()
        space(p, after=0)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    space(p, after=2)
    r = p.add_run("CONDUCTEUR DE SOUTENANCE")
    font(r, size=11, bold=True, color=GOLD)
    r.font.letter_spacing = Pt(2) if hasattr(r.font, "letter_spacing") else None

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    space(p, before=6, after=4)
    r = p.add_run("Serveur MCP intelligent\nde renseignement sur les menaces")
    font(r, size=25, bold=True, color=NAVY)
    p.paragraph_format.line_spacing = 1.05

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    space(p, after=10)
    borders(p, bottom=True, color=GOLD_HEX, size=12)
    r = p.add_run("Art de protection du secret · Cyber OPS")
    font(r, size=12.5, color=GREY)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    space(p, before=10, after=2)
    r = p.add_run("16 diapositives · 15 min 25 d'exposé · 3 questions par diapositive")
    font(r, size=11, bold=True, color=NAVY)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    space(p, after=26)
    r = p.add_run("Document de préparation — ne se lit pas à voix haute")
    font(r, size=9.5, italic=True, color=GREY)

    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.columns[0].width = Cm(15.2)
    cell = t.rows[0].cells[0]
    cell.width = Cm(15.2)
    shade(cell, PALE)
    cell.paragraphs[0].text = ""
    p = cell.paragraphs[0]
    space(p, before=2, after=2)
    r = p.add_run("Comment utiliser ce conducteur")
    font(r, size=10.5, bold=True, color=NAVY)
    for line in [
        "Une page par diapositive. En haut : la durée cible, la section et le titre.",
        "« À dire » donne les phrases prêtes, une idée par ligne ; les nombres en gras sont ceux à appuyer.",
        "« Chorégraphie » indique le geste, le second et l'élément à désigner sur la projection.",
        "« Ce que ça prouve » et « clause visée » servent si le jury demande pourquoi la diapositive existe.",
        "« Questions probables » propose des réponses de 30 secondes, avec un niveau de confiance.",
    ]:
        p = cell.add_paragraph()
        space(p, after=1)
        p.paragraph_format.left_indent = Cm(0.35)
        r = p.add_run("• ")
        font(r, size=9.2, color=GOLD)
        rich(p, line, size=9.2)

    p = doc.add_paragraph()
    space(p, before=18, after=0)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Mzabi Sirine  ·  Khaled Ellif")
    font(r, size=12, bold=True, color=NAVY)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    space(p, after=2, before=3)
    r = p.add_run("Encadrant académique : Tarek Hdji")
    font(r, size=9.4, color=GREY)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    space(p, after=8, before=0)
    r = p.add_run("TEK-UP — Université privée · Spécialité SSIR · Année universitaire 2026–2027")
    font(r, size=9, color=GREY)

    p = doc.add_paragraph()
    space(p, before=24, after=0)

    cards = [
        ("PARTIE I — SECRET", "P0 à P3, puis la clôture", "8 diapositives · Sirine"),
        ("PARTIE II — CYBER OPS", "P4 et P5", "8 diapositives · Khaled"),
        ("PREUVES", "captures, tests, journal", "7 captures réelles · 31 tests"),
    ]
    t = doc.add_table(rows=1, cols=3)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    fixed_table(t, [Cm(5.5), Cm(5.5), Cm(5.5)])
    for cell, (label, line1, line2) in zip(t.rows[0].cells, cards):
        shade(cell, PALE)
        cell.text = ""
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        space(p, before=4, after=3)
        r = p.add_run(label)
        font(r, size=7.4, bold=True, color=GOLD)
        p = cell.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        space(p, before=0, after=1)
        r = p.add_run(line1)
        font(r, size=9.4, bold=True, color=NAVY)
        p = cell.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        space(p, before=0, after=4)
        r = p.add_run(line2)
        font(r, size=8.2, color=GREY)

    doc.add_page_break()


def mode_emploi(doc):
    p = doc.add_paragraph()
    space(p, before=0, after=4)
    r = p.add_run("Vue d'ensemble et répartition de parole")
    font(r, size=17, bold=True, color=NAVY)
    borders(p, bottom=True, color=GOLD_HEX, size=10)

    p = doc.add_paragraph()
    space(p, before=6, after=8)
    rich(p, "L'exposé tient en 15 min 25. Sirine couvre la partie Art de protection du secret, "
            "de la diapositive 1 à la 7, puis conclut en 16 ; Khaled couvre Cyber OPS, de la "
            "diapositive 8 à la 15. Le passage de relais se fait à la diapositive 8, sans mot "
            "d'excuse : la phrase de liaison est écrite sur la page.", size=9.8)

    rows = [
        ("#", "Section", "Titre court", "Durée", "Voix"),
        ("1", "Ouverture", "Titre et promesse", "0:45", "Sirine"),
        ("2", "P0", "Problème réel et thèse", "1:00", "Sirine"),
        ("3", "P0", "Objectifs et périmètre", "0:55", "Sirine"),
        ("4", "P1", "Menaces et risques classés", "1:00", "Sirine"),
        ("5", "P1", "Actifs et étiquettes", "0:55", "Sirine"),
        ("6", "P1", "Exigences EF / ENF", "1:00", "Sirine"),
        ("7", "P2/P3", "Comparatif et architecture", "1:00", "Sirine"),
        ("8", "P3", "Zones et points de contrôle", "1:10", "Khaled"),
        ("9", "P4", "Réalisation et journal", "1:00", "Khaled"),
        ("10", "P4", "Durcissement", "0:55", "Khaled"),
        ("11", "P5", "Tests : la fuite bloquée", "0:50", "Khaled"),
        ("12", "P5", "Attaques et résultats", "1:10", "Khaled"),
        ("13", "P5", "Métriques et avant/après", "1:10", "Khaled"),
        ("14", "P5", "Honnêteté des statuts", "0:55", "Khaled"),
        ("15", "Bilan", "Limites et travail restant", "0:50", "Khaled"),
        ("16", "Clôture", "Conclusion et ouverture", "0:50", "Sirine"),
    ]
    t = doc.add_table(rows=len(rows), cols=5)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    widths = [Cm(1.0), Cm(2.2), Cm(7.9), Cm(2.1), Cm(4.2)]
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            cell = t.rows[i].cells[j]
            cell.text = ""
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if j in (0, 3, 4) else WD_ALIGN_PARAGRAPH.LEFT
            space(p, before=1, after=1)
            if i == 0:
                shade(cell, NAVY_HEX)
                r = p.add_run(val)
                font(r, size=9.0, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
            else:
                if i % 2 == 0:
                    shade(cell, PALE)
                r = p.add_run(val)
                font(r, size=9.0, bold=(j in (0, 3)), color=NAVY if j == 0 else INK)
    fixed_table(t, widths)
    t.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))

    p = doc.add_paragraph()
    space(p, before=10, after=4)
    r = p.add_run("Deux règles de conduite")
    font(r, size=11.5, bold=True, color=NAVY)
    for label, text in [
        ("Regard", "on annonce ce que le jury va voir, on le montre, puis on dit ce que cela prouve. Jamais de commentaire pendant que la machine travaille."),
        ("Preuve", "chaque affirmation renvoie à une capture, un test ou un chiffre du rapport ; sinon la formule est « proposé » ou « à vérifier »."),
    ]:
        p = doc.add_paragraph()
        space(p, before=2, after=2)
        p.paragraph_format.left_indent = Cm(0.3)
        r = p.add_run(f"{label} — ")
        font(r, size=9.6, bold=True, color=GOLD)
        rich(p, text, size=9.6)

    p = doc.add_paragraph()
    space(p, before=8, after=2)
    rich(p, "Le total des durées affichées est de 925 secondes. Garder 35 secondes de marge "
            "sur les transitions : viser 14 min 50 à la montre, pour finir avant la limite.",
         size=9.2, color=GREY)

    p = doc.add_paragraph()
    space(p, before=14, after=4)
    r = p.add_run("Checklist avant d'entrer")
    font(r, size=11.5, bold=True, color=NAVY)

    checks = [
        ("Poste", "serveur lancé sur une base vérifiée ; terminal agrandi ; police lisible."),
        ("Support", "deck PDF et rapport PDF sur le bureau et sur clé USB ; dossier preuves/ ouvert."),
        ("Réseau", "mode avion ou Wi-Fi coupé ; aucun service externe requis par la démonstration."),
        ("Temps", "minuteur visible à 15 min ; le passage de relais est annoncé à la diapositive 8."),
        ("Voix", "une phrase par idée ; les nombres sont dits lentement, sans lire la projection."),
        ("Repli", "si la machine hésite, montrer la capture correspondante et continuer sans commenter."),
    ]
    t = doc.add_table(rows=3, cols=2)
    fixed_table(t, [Cm(8.7), Cm(8.7)])
    for i, (label, text) in enumerate(checks):
        cell = t.rows[i // 2].cells[i % 2]
        shade(cell, PALE)
        cell.text = ""
        p = cell.paragraphs[0]
        space(p, before=2, after=2)
        r = p.add_run(f"□  {label} — ")
        font(r, size=9.0, bold=True, color=NAVY)
        rich(p, text, size=9.0)

    doc.add_page_break()


def slide_page(doc, s):
    # Bandeau : numéro, section, durée
    p = doc.add_paragraph()
    space(p, before=0, after=2)
    tab = p.paragraph_format.tab_stops.add_tab_stop(Cm(17.4), WD_ALIGN_PARAGRAPH.RIGHT)
    r = p.add_run(f"DIAPOSITIVE {s['n']:02d}   ·   {s['section']}")
    font(r, size=9.2, bold=True, color=GOLD)
    r = p.add_run(f"\t{s['duration']}")
    font(r, size=9.2, bold=True, color=NAVY)

    p = doc.add_paragraph()
    space(p, before=1, after=4)
    r = p.add_run(s["title"])
    font(r, size=18, bold=True, color=NAVY)
    borders(p, bottom=True, color=GOLD_HEX, size=6)

    # À dire
    p = doc.add_paragraph()
    space(p, before=4, after=2)
    r = p.add_run("À DIRE")
    font(r, size=9.0, bold=True, color=GOLD)
    for i, line in enumerate(s["say"], 1):
        p = doc.add_paragraph()
        space(p, before=0, after=2.0, line=1.03)
        p.paragraph_format.left_indent = Cm(0.55)
        p.paragraph_format.first_line_indent = Cm(-0.55)
        r = p.add_run(f"{i}. ")
        font(r, size=10.8, bold=True, color=GOLD)
        rich(p, line, size=10.8)

    # Preuve + clause
    t = doc.add_table(rows=1, cols=2)
    t.autofit = False
    fixed_table(t, [Cm(8.4), Cm(9.0)])
    for j, (label, text) in enumerate((
        ("CE QUE ÇA PROUVE", s["proof"]),
        ("CLAUSE VISÉE", s["clause"]),
    )):
        cell = t.rows[0].cells[j]
        cell.width = Cm(8.4 if j == 0 else 9.0)
        shade(cell, PALE)
        cell.text = ""
        p = cell.paragraphs[0]
        space(p, before=1, after=0)
        r = p.add_run(label)
        font(r, size=8.0, bold=True, color=GREY)
        p = cell.add_paragraph()
        space(p, before=0, after=1)
        rich(p, text, size=9.6, bold_numbers=True)

    # Chorégraphie + Q&R côte à côte
    p = doc.add_paragraph()
    space(p, before=5, after=2)
    r = p.add_run("CHORÉGRAPHIE")
    font(r, size=9.0, bold=True, color=GOLD)

    for line in s["choreo"]:
        p = doc.add_paragraph()
        space(p, before=0, after=0.8, line=1.02)
        p.paragraph_format.left_indent = Cm(0.4)
        p.paragraph_format.first_line_indent = Cm(-0.4)
        r = p.add_run("→ ")
        font(r, size=10.0, color=GOLD)
        rich(p, line, size=10.0)

    p = doc.add_paragraph()
    space(p, before=5, after=2)
    r = p.add_run("QUESTIONS PROBABLES DU JURY")
    font(r, size=9.0, bold=True, color=GOLD)

    for q, a, conf in s["qa"]:
        p = doc.add_paragraph()
        space(p, before=2, after=0.5)
        r = p.add_run(q)
        font(r, size=10.4, bold=True, color=NAVY)
        r = p.add_run(f"   [confiance {conf}]")
        font(r, size=7.6, italic=True, color=GREY)
        p = doc.add_paragraph()
        space(p, before=0, after=1.5, line=1.02)
        p.paragraph_format.left_indent = Cm(0.35)
        rich(p, a, size=10.0)

    # Repli + prononciation
    t = doc.add_table(rows=1, cols=1)
    t.autofit = False
    fixed_table(t, [Cm(17.4)])
    cell = t.rows[0].cells[0]
    shade(cell, PALE_GOLD)
    cell.text = ""
    p = cell.paragraphs[0]
    space(p, before=1, after=1)
    r = p.add_run("REPLI 30 SECONDES  ")
    font(r, size=8.2, bold=True, color=GOLD)
    rich(p, s["fallback"], size=10.0)

    p = doc.add_paragraph()
    space(p, before=3, after=0)
    r = p.add_run("Prononciation : ")
    font(r, size=8.4, bold=True, color=GREY)
    rich(p, s["pron"], size=8.4, color=GREY, bold_numbers=False)

    p = doc.add_paragraph()
    space(p, before=10, after=2)
    r = p.add_run("NOTES DE SÉANCE")
    font(r, size=7.4, bold=True, color=GREY)
    for _ in range(4):
        p = doc.add_paragraph()
        space(p, before=13, after=0)
        p.paragraph_format.tab_stops.add_tab_stop(
            Cm(17.4), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.LINES)
        r = p.add_run("\t")
        font(r, size=8.4, color=RGBColor(0xC2, 0xCB, 0xD6))

    doc.add_page_break()


def secours(doc):
    p = doc.add_paragraph()
    space(p, before=0, after=4)
    r = p.add_run("Plan de secours — version 10 minutes")
    font(r, size=17, bold=True, color=NAVY)
    borders(p, bottom=True, color=GOLD_HEX, size=10)

    p = doc.add_paragraph()
    space(p, before=6, after=8)
    rich(p, "Si le jury réduit le temps, garder le fil « problème → preuve → limite ». "
            "Les diapositives 3, 5 et 9 se résument en une phrase orale et se sautent sans "
            "casser le raisonnement. Le total ci-dessous fait 9 min 30, laissant 30 secondes "
            "de respiration.", size=9.8)

    rows = [
        ("#", "Diapositive", "Secondes", "Ce qu'on ne coupe pas"),
        ("1", "Titre et promesse", "30", "les deux validations et le fil rouge"),
        ("2", "Problème et thèse", "40", "le chiffre du problème et la phrase de thèse"),
        ("4", "Menaces et risques", "40", "la matrice et les deux risques majeurs"),
        ("6", "Exigences", "40", "EF de confidentialité et ENF d'audit"),
        ("7", "Comparatif et architecture", "40", "le choix C4 et les trois zones"),
        ("8", "Zones et CP", "60", "CP1, CP2, CP5, CP7 et le journal unique"),
        ("10", "Durcissement", "40", "uid, droits de la base, zéro port"),
        ("11", "Fuite bloquée", "40", "les deux refus et les deux motifs"),
        ("12", "Attaques", "70", "16/16 abus et 3/3 injections"),
        ("13", "Métriques", "50", "8/8 contre 0/8 et la p95"),
        ("14", "Statuts", "50", "testé / configuré / proposé"),
        ("15", "Limites", "40", "les cinq limites assumées"),
        ("16", "Conclusion", "30", "la phrase de clôture"),
    ]
    t = doc.add_table(rows=len(rows), cols=4)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    widths = [Cm(1.0), Cm(6.2), Cm(2.0), Cm(8.2)]
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            cell = t.rows[i].cells[j]
            cell.text = ""
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if j in (0, 2) else WD_ALIGN_PARAGRAPH.LEFT
            space(p, before=1, after=1)
            if i == 0:
                shade(cell, NAVY_HEX)
                r = p.add_run(val)
                font(r, size=8.8, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
            else:
                if i % 2 == 0:
                    shade(cell, PALE)
                r = p.add_run(val)
                font(r, size=8.8, bold=(j in (0, 2)), color=NAVY if j == 0 else INK)
    fixed_table(t, widths)
    t.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))

    p = doc.add_paragraph()
    space(p, before=10, after=3)
    r = p.add_run("Phrases de transition à ne pas perdre")
    font(r, size=11.5, bold=True, color=NAVY)
    for text in [
        "Après la diapositive 7 : « La décision est prise ; il reste à prouver qu'elle tient. Khaled prend la suite sur la réalisation. »",
        "Avant la démonstration : « Je ne raconte pas le contrôle, je le déclenche devant vous. »",
        "Après la démonstration : « La fuite est bloquée, journalisée et visible : c'est la promesse tenue, pas une intention. »",
        "Conclusion : « Le secret n'est pas une promesse de politique, c'est une propriété testée ; les limites restantes sont nommées. »",
    ]:
        p = doc.add_paragraph()
        space(p, before=2, after=2)
        p.paragraph_format.left_indent = Cm(0.4)
        p.paragraph_format.first_line_indent = Cm(-0.4)
        r = p.add_run("— ")
        font(r, size=9.4, color=GOLD)
        rich(p, text, size=9.2)

    p = doc.add_paragraph()
    space(p, before=10, after=2)
    r = p.add_run("Si un test échoue pendant la soutenance")
    font(r, size=11.5, bold=True, color=NAVY)
    for text in [
        "Annoncer le statut réel sans improviser : « le test est rejoué maintenant ; s'il échoue, je montre le journal et la correction prévue ».",
        "Ne jamais masquer une sortie : commenter le code obtenu, puis revenir au résultat déjà consigné dans preuves/logs/tests_latest.json.",
        "Si la machine ne répond plus, basculer sur les captures 01 à 07 de preuves/captures_ecran : elles portent l'horodatage et le numéro de commande.",
    ]:
        p = doc.add_paragraph()
        space(p, before=1, after=1)
        p.paragraph_format.left_indent = Cm(0.4)
        p.paragraph_format.first_line_indent = Cm(-0.4)
        r = p.add_run("• ")
        font(r, size=9.2, color=GOLD)
        rich(p, text, size=9.2)

    p = doc.add_paragraph()
    space(p, before=12, after=0)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Fin du conducteur · 16 diapositives · dossier preuves/ joint sur clé USB")
    font(r, size=8.2, italic=True, color=GREY)


def build():
    doc = Document()
    page_setup(doc)

    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(9.4)
    style.paragraph_format.space_after = Pt(0)

    header_footer(doc)
    settings = doc.settings.element
    upd = OxmlElement("w:updateFields")
    upd.set(qn("w:val"), "true")
    settings.append(upd)

    cover(doc)
    mode_emploi(doc)
    for s in SLIDES:
        slide_page(doc, s)
    secours(doc)

    doc.save(OUT)
    print(f"[docx] {OUT}")

    subprocess.run(
        ["libreoffice", "--headless", "--convert-to", "pdf", "--outdir", str(OUT.parent), str(OUT)],
        check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL,
    )
    pdf = OUT.with_suffix(".pdf")
    print(f"[pdf ] {pdf}")
    info = subprocess.run(["pdfinfo", str(pdf)], capture_output=True, text=True).stdout
    for line in info.splitlines():
        if line.startswith(("Pages:", "Page size:")):
            print("      ", line)


if __name__ == "__main__":
    build()
