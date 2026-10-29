#!/usr/bin/env python3
"""Construit le deck de soutenance 16:9 et ses notes orales."""
from __future__ import annotations

import json
import re
import subprocess
import sys
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

sys.path.insert(0, str(Path(__file__).resolve().parent))
from deck_content import MEMOS, SLIDES  # noqa: E402

PROJECT = Path(__file__).resolve().parents[2]
FIG = PROJECT / "figures"
CAP = PROJECT / "preuves" / "captures"
CAP_REEL = PROJECT / "preuves" / "captures_reelles"
CAP_ECRAN = PROJECT / "preuves" / "captures_ecran"
OUT = Path(__file__).resolve().parent / "Deck_Soutenance_MCP_TI.pptx"

FONT = "Arial"
MONO = "Consolas"
NAVY = RGBColor(0x0B, 0x1F, 0x3A)
INK = RGBColor(0x18, 0x26, 0x3A)
BLUE = RGBColor(0x1D, 0x5B, 0x94)
GOLD = RGBColor(0xB8, 0x93, 0x4E)
RED = RGBColor(0xA6, 0x3B, 0x3B)
GREEN = RGBColor(0x2D, 0x7A, 0x5B)
MUTED = RGBColor(0x6D, 0x7B, 0x8D)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
PAPER = RGBColor(0xFB, 0xFC, 0xFE)
LIGHT = RGBColor(0xF2, 0xF5, 0xF9)
PALE_BLUE = RGBColor(0xE7, 0xEF, 0xF8)
PALE_GOLD = RGBColor(0xF7, 0xF1, 0xE3)
PALE_RED = RGBColor(0xF8, 0xEA, 0xEA)
PALE_GREEN = RGBColor(0xE7, 0xF3, 0xED)
LINE = RGBColor(0xD9, 0xE1, 0xEA)
DARK_TEXT = RGBColor(0x1F, 0x2A, 0x3A)

W, H = 13.333, 7.5
M = 0.68
CONTENT_Y = 1.97


def rgb(hex_value: str) -> RGBColor:
    return RGBColor.from_string(hex_value.replace("#", ""))


def rect(slide, x, y, w, h, fill, line=None, shape=MSO_SHAPE.RECTANGLE, radius=None):
    shp = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = Pt(0.8)
    if radius is not None and shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        shp.adjustments[0] = radius
    return shp


def _apply_text(p, text, size, bold, color, font, italic=False):
    for token in re.split(r"(\*\*|__|`[^`]+`)", text):
        if token == "**":
            bold = not bold
            continue
        if token == "__":
            italic = not italic
            continue
        if not token:
            continue
        if token.startswith("`") and token.endswith("`"):
            run = p.add_run()
            run.text = token[1:-1]
            run.font.name = MONO
            run.font.size = Pt(max(8, size - 0.6))
            run.font.bold = bold
            run.font.italic = italic
            run.font.color.rgb = color
            continue
        run = p.add_run()
        run.text = token
        run.font.name = font
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.italic = italic
        run.font.color.rgb = color


def text(slide, x, y, w, h, value, size=14, bold=False, color=INK, font=FONT,
         align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, italic=False,
         line_spacing=1.0, space_after=0, margins=(0, 0, 0, 0)):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left, tf.margin_right, tf.margin_top, tf.margin_bottom = [Inches(v) for v in margins]
    tf.vertical_anchor = anchor
    lines = str(value).split("\n")
    for index, line in enumerate(lines):
        p = tf.paragraphs[0] if index == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = line_spacing
        p.space_after = Pt(space_after)
        _apply_text(p, line, size, bold, color, font, italic)
    return box


def title_slide(prs, s):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    rect(slide, 0, 0, W, H, NAVY)
    rect(slide, M, 1.18, 0.07, 3.55, GOLD)
    text(slide, 1.02, 1.08, 8.0, 0.35, "PROJET ACADÉMIQUE · SSIR · 2026–2027",
         size=11, bold=True, color=GOLD)
    text(slide, 1.02, 1.47, 8.5, 2.0,
         "Serveur MCP intelligent de renseignement sur les menaces",
         size=31, bold=True, color=WHITE, line_spacing=1.02)
    text(slide, 1.02, 3.52, 8.4, 0.75,
         "Protection du secret, étiquetage et traçabilité de bout en bout",
         size=15.5, italic=True, color=GOLD)
    text(slide, 1.02, 4.55, 5.2, 0.4, "Mzabi Sirine · Khaled Ellif",
         size=20, bold=True, color=WHITE)
    text(slide, 1.02, 5.0, 7.6, 0.35,
         "Encadrant académique : Tarek Hdji",
         size=10.5, color=RGBColor(0xC8, 0xD3, 0xE0))
    chip(slide, 1.02, 5.62, 2.86, 0.44, "VALIDATION 1 · ARCHITECTURE", RGBColor(0x16, 0x2F, 0x52), RGBColor(0xE6, 0xD6, 0xB2))
    chip(slide, 4.02, 5.62, 3.35, 0.44, "VALIDATION 2 · RÉALISATION & TESTS", RGBColor(0x16, 0x2F, 0x52), RGBColor(0xE6, 0xD6, 0xB2))
    rect(slide, 1.02, 6.38, 11.3, 0.015, RGBColor(0x3B, 0x52, 0x70))
    text(slide, 1.02, 6.55, 8.9, 0.3,
         "Version 1.0 — 27 octobre 2026 · données synthétiques · laboratoire hors ligne",
         size=9.5, color=RGBColor(0xAF, 0xBE, 0xD0))
    text(slide, 9.5, 6.55, 2.9, 0.3, "TEK-UP · SSIR", size=9.5, bold=True,
         color=GOLD, align=PP_ALIGN.RIGHT)
    rect(slide, 10.25, 1.35, 2.55, 1.45, WHITE, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.08)
    add_image(slide, FIG / "logoTek-up_university.png", 10.32, 1.45, 2.4, 1.25)
    text(slide, 10.18, 2.95, 3.0, 0.3, "UNIVERSITÉ PRIVÉE TEK-UP",
         size=8.5, bold=True, color=RGBColor(0xD7, 0xC1, 0x92), align=PP_ALIGN.CENTER)
    notes(slide, s)


def chrome(prs, s, dark=False):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    if dark:
        rect(slide, 0, 0, W, H, NAVY)
    else:
        rect(slide, 0, 0, W, H, PAPER)
        rect(slide, 0, 0, W, 0.055, NAVY)
    kicker_color = GOLD if dark else GOLD
    title_color = WHITE if dark else NAVY
    text(slide, M, 0.38, 10.7, 0.3, s["section"], size=10, bold=True, color=kicker_color)
    size = 26 if len(s["title"]) < 45 else (24 if len(s["title"]) < 60 else 21.5)
    text(slide, M, 0.7, 11.7, 0.95, s["title"], size=size, bold=True,
         color=title_color, line_spacing=1.02)
    text(slide, 10.1, 0.38, 2.55, 0.3, "TEK-UP · SSIR", size=8.5, bold=True,
         color=(RGBColor(0xB9, 0xC7, 0xD8) if dark else MUTED), align=PP_ALIGN.RIGHT)
    rect(slide, M, 1.73, 11.95, 0.018, GOLD)
    footer(slide, s["n"], dark=dark)
    return slide


def footer(slide, n, dark=False):
    color = RGBColor(0x9F, 0xB0, 0xC4) if dark else MUTED
    rect(slide, M, 7.08, 11.95, 0.012, RGBColor(0x3B, 0x52, 0x70) if dark else LINE)
    text(slide, M, 7.17, 8.4, 0.22, "Serveur MCP de renseignement sur les menaces",
         size=8, color=color)
    text(slide, 10.5, 7.17, 2.1, 0.22, f"{n:02d} / 16", size=8.5, bold=True,
         color=color, align=PP_ALIGN.RIGHT)


def chip(slide, x, y, w, h, label, fill, fg=WHITE, size=8.1, bold=True):
    rect(slide, x, y, w, h, fill, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.18)
    text(slide, x + 0.09, y + 0.02, w - 0.18, h - 0.04, label, size=size, bold=bold,
         color=fg, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)


def card(slide, x, y, w, h, fill=LIGHT, border=None, radius=0.025):
    return rect(slide, x, y, w, h, fill, line=border,
                shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=radius)


def kicker_number(slide, x, y, w, h, value, label, sub, accent=BLUE, size=31):
    card(slide, x, y, w, h, LIGHT)
    rect(slide, x, y, 0.06, h, accent)
    text(slide, x + 0.2, y + 0.12, w - 0.3, 0.55, value, size=size, bold=True, color=accent)
    text(slide, x + 0.2, y + 0.69, w - 0.3, 0.3, label, size=10.5, bold=True, color=INK)
    text(slide, x + 0.2, y + 0.98, w - 0.3, 0.32, sub, size=8.3, color=MUTED)


def add_image(slide, path: Path, x, y, w, h):
    with Image.open(path) as im:
        iw, ih = im.size
    ratio = iw / ih
    box_ratio = w / h
    if ratio >= box_ratio:
        draw_w = w
        draw_h = w / ratio
    else:
        draw_h = h
        draw_w = h * ratio
    dx = x + (w - draw_w) / 2
    dy = y + (h - draw_h) / 2
    slide.shapes.add_picture(str(path), Inches(dx), Inches(dy), Inches(draw_w), Inches(draw_h))


def table(slide, x, y, w, h, headers, rows, widths, font_size=11.2, header_size=10.4):
    shape = slide.shapes.add_table(len(rows) + 1, len(headers), Inches(x), Inches(y), Inches(w), Inches(h))
    tbl = shape.table
    for index, col_w in enumerate(widths):
        tbl.columns[index].width = Inches(col_w)
    tbl.rows[0].height = Inches(0.42)
    for r_index in range(1, len(rows) + 1):
        tbl.rows[r_index].height = Inches((h - 0.42) / len(rows))
    for c_index, label in enumerate(headers):
        cell = tbl.cell(0, c_index)
        cell.fill.solid()
        cell.fill.fore_color.rgb = NAVY
        cell.margin_left = Inches(0.08)
        cell.margin_right = Inches(0.06)
        cell.margin_top = Inches(0.02)
        cell.margin_bottom = Inches(0.02)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf = cell.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER if c_index in (0, 1) else PP_ALIGN.LEFT
        _apply_text(p, label, header_size, True, WHITE, FONT)
    for r_index, row in enumerate(rows, start=1):
        for c_index, value in enumerate(row):
            cell = tbl.cell(r_index, c_index)
            cell.fill.solid()
            cell.fill.fore_color.rgb = WHITE if r_index % 2 else LIGHT
            cell.margin_left = Inches(0.08)
            cell.margin_right = Inches(0.06)
            cell.margin_top = Inches(0.025)
            cell.margin_bottom = Inches(0.025)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.CENTER if c_index in (0, 1) else PP_ALIGN.LEFT
            _apply_text(p, str(value), font_size, c_index == 0, INK, FONT)
    return shape


def notes(slide, s):
    tf = slide.notes_slide.notes_text_frame
    tf.clear()
    blocks = [
        ("MÉMO À DIRE — " + s["duration"], "\n".join("• " + line for line in MEMOS.get(s["n"], []))),
        ("À DIRE — " + s["duration"], None),
        ("\n".join("• " + line for line in s["say"]), None),
        ("CE QUE ÇA PROUVE", s["proof"]),
        ("CLAUSE VISÉE", s["clause"]),
        ("CHORÉGRAPHIE", "\n".join("• " + line for line in s["choreo"])),
        ("QUESTIONS PROBABLES", "\n".join(
            f"Q{i} — {q}\nRéponse 30 s : {a}\nConfiance : {conf}."
            for i, (q, a, conf) in enumerate(s["qa"], start=1))),
        ("REPLI 30 SECONDES", s["fallback"]),
        ("PRONONCIATION", s["pron"]),
    ]
    first = True
    for heading, body in blocks:
        if first:
            p = tf.paragraphs[0]
            first = False
        else:
            p = tf.add_paragraph()
        p.space_before = Pt(5)
        p.space_after = Pt(2)
        r = p.add_run()
        r.text = heading
        r.font.name = FONT
        r.font.size = Pt(10)
        r.font.bold = True
        r.font.color.rgb = NAVY
        if body:
            for line in body.split("\n"):
                p2 = tf.add_paragraph()
                p2.space_after = Pt(1)
                p2.line_spacing = 1.0
                _apply_text(p2, line, 9.5, False, DARK_TEXT, FONT)


def slide_02(prs, s):
    slide = chrome(prs, s)
    card(slide, M, CONTENT_Y, 4.6, 2.3, PALE_GOLD, border=GOLD)
    text(slide, M + 0.25, CONTENT_Y + 0.18, 4.1, 0.25, "LA DEMANDE QUI SEMBLE ANODINE",
         size=9, bold=True, color=GOLD)
    text(slide, M + 0.25, CONTENT_Y + 0.56, 4.0, 0.7, "« Résume\nADV-2026-004. »", size=28,
         bold=True, color=NAVY, line_spacing=1.0)
    text(slide, M + 0.25, CONTENT_Y + 1.84, 4.0, 0.3, "Un seul appel. Trois secrets possibles.",
         size=10.5, italic=True, color=MUTED)
    items = [
        ("AGENT", "Il exécute la demande… et une instruction cachée.", RED),
        ("BULLETIN", "Secret : canari de fuite + identité de source.", GOLD),
        ("RÉPONSE", "Trop riche = fuite ; trop vague = agent inutile.", BLUE),
    ]
    for index, (label, body, accent) in enumerate(items):
        yy = CONTENT_Y + index * 1.08
        card(slide, 5.65, yy, 6.98, 0.93)
        rect(slide, 5.65, yy, 0.06, 0.93, accent)
        text(slide, 5.9, yy + 0.12, 1.1, 0.26, label, size=9.5, bold=True, color=accent)
        text(slide, 7.05, yy + 0.12, 5.35, 0.55, body, size=13.2, color=INK)
    stats = [
        ("4,44 M$", "coût moyen d'une violation", "IBM 2025", RED),
        ("12 / 30 / 6", "bulletins / IOC / actifs", "jeu synthétique", BLUE),
        ("31", "scénarios de campagne", "15 nominaux + 16 abus", GREEN),
    ]
    for index, (value, label, sub, accent) in enumerate(stats):
        x = 0.68 + index * 3.97
        kicker_number(slide, x, 4.85, 3.75, 1.55, value, label, sub, accent=accent, size=26)
    notes(slide, s)


def slide_03(prs, s):
    slide = chrome(prs, s)
    objectives = [
        ("O1", "0 fuite", "16 / 16 abus", RED),
        ("O2", "100 % tracé", "115 / 115 appels", BLUE),
        ("O3", "3 / 3 inject.", "0 faux positif", GREEN),
        ("O4", "p95 ≤ 50 ms", "12,42 ms", BLUE),
        ("O5", "≤ 5 s", "10,0 ms · uid 65534", GOLD),
    ]
    for index, (oid, value, result, accent) in enumerate(objectives):
        x = 0.68 + index * 2.42
        card(slide, x, CONTENT_Y, 2.22, 2.55, LIGHT)
        text(slide, x + 0.17, CONTENT_Y + 0.15, 0.6, 0.28, oid, size=11, bold=True, color=accent)
        text(slide, x + 0.17, CONTENT_Y + 0.55, 1.9, 0.5, value, size=17.5, bold=True, color=NAVY)
        rect(slide, x + 0.17, CONTENT_Y + 1.18, 1.5, 0.018, accent)
        text(slide, x + 0.17, CONTENT_Y + 1.38, 1.9, 0.55, result, size=10.5, color=MUTED)
    card(slide, M, 4.85, 5.75, 1.55, PALE_GREEN, border=GREEN)
    text(slide, M + 0.22, 4.98, 5.2, 0.25, "INCLUS DANS LA MAQUETTE", size=9, bold=True, color=GREEN)
    text(slide, M + 0.22, 5.3, 5.25, 0.95,
         "stdio · 6 outils · ABAC · étiquetage · audit chaîné · « nobody » · base 0440 · clé 0400",
         size=10.6, color=INK, line_spacing=1.12)
    card(slide, 6.72, 4.85, 5.92, 1.55, PALE_RED, border=RED)
    text(slide, 6.94, 4.98, 5.4, 0.25, "EXCLU OU LAISSÉ À L'INTÉGRATION", size=9, bold=True, color=RED)
    text(slide, 6.94, 5.3, 5.4, 0.95,
         "HTTP · OAuth 2.1 · mTLS · EDR réel · ancrage externe du journal · données réelles",
         size=10.6, color=INK, line_spacing=1.12)
    notes(slide, s)


def slide_04(prs, s):
    slide = chrome(prs, s)
    cols = [
        ("ACTIFS", "Ce qu'on protège", [
            ("Secret", "source, opération CANARI"),
            ("Confidentiel", "vulnérabilités, actifs"),
            ("Interne", "indicateurs, familles"),
            ("Public", "bulletins de sensibilisation"),
            ("Secrets op.", "clé, base, journal"),
        ], GOLD),
        ("ACTEURS", "Qui peut agir", [
            ("Analyste SOC", "habilitation 1–2"),
            ("Lead CTI", "habilitation 3 + compartiments"),
            ("Agent conversationnel", "légitime mais manipulable"),
            ("Partenaire", "fournit, ne commande pas"),
            ("Attaquant", "jeton, flux, poste"),
        ], BLUE),
        ("SURFACE", "Où l'on peut frapper", [
            ("stdio", "canal local"),
            ("6 outils", "schémas stricts"),
            ("Ressources", "URI filtrées"),
            ("Contenu importé", "injection, poison"),
            ("Fichiers", "clé, base, journal"),
        ], RED),
    ]
    for index, (heading, sub, items, accent) in enumerate(cols):
        x = M + index * 4.15
        width = 3.95 if index < 2 else 3.7
        card(slide, x, CONTENT_Y, width, 3.95, LIGHT)
        text(slide, x + 0.2, CONTENT_Y + 0.18, width - 0.4, 0.28, heading, size=9.5, bold=True, color=accent)
        text(slide, x + 0.2, CONTENT_Y + 0.48, width - 0.4, 0.3, sub, size=12.5, bold=True, color=NAVY)
        yy = CONTENT_Y + 0.95
        for label, detail in items:
            text(slide, x + 0.2, yy, 0.25, 0.25, "●", size=7.5, color=accent)
            text(slide, x + 0.43, yy - 0.005, width - 0.6, 0.24, label, size=10.3, bold=True, color=INK)
            text(slide, x + 0.43, yy + 0.23, width - 0.6, 0.35, detail, size=8.8, color=MUTED)
            yy += 0.56
    card(slide, M, 6.12, 11.95, 0.68, PALE_GOLD, border=GOLD)
    text(slide, M + 0.22, 6.22, 11.5, 0.45,
         "L'agent est légitime, mais il exécute aussi une instruction cachée : le serveur ne traite jamais le contenu comme une commande.",
         size=11.2, bold=True, color=NAVY)
    notes(slide, s)


def slide_05(prs, s):
    slide = chrome(prs, s)
    add_image(slide, FIG / "fig_matrice_risques.png", 0.45, CONTENT_Y, 8.05, 4.72)
    text(slide, 0.55, 6.75, 7.9, 0.22,
         "Source : modèle de menace du projet — scores argumentés, pas des mesures statistiques.",
         size=7.8, italic=True, color=MUTED)
    top = [
        ("R01", "Fuite Secret", "20", RED, "réponse trop riche"),
        ("R02", "Injection", "20", RED, "flux partenaire"),
        ("R04", "Hors compartiment", "16", GOLD, "besoin d'en connaître"),
    ]
    for index, (rid, label, score, accent, detail) in enumerate(top):
        yy = CONTENT_Y + index * 1.16
        card(slide, 8.85, yy, 3.85, 1.0, LIGHT)
        rect(slide, 8.85, yy, 0.06, 1.0, accent)
        text(slide, 9.08, yy + 0.12, 0.6, 0.3, rid, size=10, bold=True, color=accent)
        text(slide, 9.74, yy + 0.11, 1.9, 0.3, label, size=12.5, bold=True, color=NAVY)
        text(slide, 12.0, yy + 0.02, 0.6, 0.55, score, size=22, bold=True, color=accent, align=PP_ALIGN.RIGHT)
        text(slide, 9.08, yy + 0.55, 3.35, 0.3, detail, size=9, color=MUTED)
    card(slide, 8.85, 5.55, 3.85, 1.1, PALE_GOLD, border=GOLD)
    text(slide, 9.05, 5.68, 3.45, 0.85,
         "Un score de 15 reste grave : clé compromise ou traces altérées ont un impact maximal.",
         size=9.8, bold=True, color=NAVY, line_spacing=1.08)
    notes(slide, s)


def slide_06(prs, s):
    slide = chrome(prs, s)
    chips = [("13 exigences fonctionnelles", BLUE), ("12 exigences non fonctionnelles", GREEN),
             ("10 risques couverts", GOLD)]
    for index, (label, color) in enumerate(chips):
        chip(slide, M + index * 4.0, 1.95, 3.7, 0.42, label, color, WHITE, size=9.2)
    rows = [
        ["R01", "Fuite d'un bulletin Secret", "EF-03 · EF-04", "CP2 · CP5", "A06 · A07 · A15"],
        ["R02", "Injection partenaire", "EF-05 · ENF-12", "CP4 · CP5", "A10"],
        ["R06", "Altération des traces", "EF-08 · ENF-08", "CP7", "A14"],
        ["R10", "Fichiers accessibles", "ENF-06 · ENF-07", "CP8 · N15", "N15"],
    ]
    table(slide, M, 2.62, 11.95, 2.8,
          ["Risque", "Ce qui peut arriver", "Exigences", "Contrôles", "Tests"],
          rows, [0.9, 4.0, 3.1, 2.3, 1.65], font_size=10.4, header_size=9.6)
    card(slide, M, 5.72, 11.95, 0.95, PALE_BLUE, border=BLUE)
    text(slide, M + 0.22, 5.84, 11.5, 0.7,
         "Un contrôle sans exigence est décoratif ; une exigence sans test est une intention. La matrice se lit dans les deux sens.",
         size=11.4, bold=True, color=NAVY, line_spacing=1.05)
    notes(slide, s)


def slide_07(prs, s):
    slide = chrome(prs, s)
    add_image(slide, FIG / "fig_comparatif.png", 0.42, CONTENT_Y, 8.1, 4.35)
    text(slide, 0.55, 6.5, 8.0, 0.22,
         "Source : notation pondérée du projet ; les poids sont fixés avant les notes.",
         size=7.8, italic=True, color=MUTED)
    card(slide, 8.82, CONTENT_Y, 3.88, 1.45, PALE_GREEN, border=GREEN)
    text(slide, 9.05, CONTENT_Y + 0.13, 2.4, 0.3, "C4 — SERVEUR DÉDIÉ", size=9, bold=True, color=GREEN)
    text(slide, 9.05, CONTENT_Y + 0.38, 1.7, 0.6, "4,18", size=33, bold=True, color=GREEN)
    text(slide, 10.35, CONTENT_Y + 0.55, 2.1, 0.35, "sur 5", size=12, color=MUTED)
    text(slide, 9.05, CONTENT_Y + 1.02, 3.4, 0.28,
         "sécurité 30 % + conformité 20 %", size=8.8, color=INK)
    lines = [
        ("C1", "2,15", "script local — le plus populaire", RED),
        ("C2", "2,68", "plateforme commerciale", MUTED),
        ("C3", "3,03", "passerelle MCP générique", BLUE),
    ]
    for index, (cid, score, label, accent) in enumerate(lines):
        yy = 3.62 + index * 0.68
        text(slide, 8.95, yy, 0.5, 0.25, cid, size=9.5, bold=True, color=accent)
        text(slide, 9.5, yy - 0.03, 0.75, 0.3, score, size=14, bold=True, color=accent)
        text(slide, 10.28, yy, 2.4, 0.25, label, size=8.6, color=MUTED)
    card(slide, 8.82, 5.68, 3.88, 0.9, PALE_RED, border=RED)
    text(slide, 9.02, 5.8, 3.45, 0.65,
         "C1 écarté : ni jeton, ni étiquette, ni journal. Immédiat ne veut pas dire sûr.",
         size=9.2, bold=True, color=NAVY, line_spacing=1.03)
    notes(slide, s)


def slide_08(prs, s):
    slide = chrome(prs, s)
    add_image(slide, FIG / "fig_architecture_zones.png", 0.38, CONTENT_Y, 8.18, 4.72)
    text(slide, 0.55, 6.75, 7.9, 0.22,
         "Source : conception du projet — CP1 à CP8 détaillés dans le rapport.",
         size=7.8, italic=True, color=MUTED)
    zones = [
        ("ZONE 0", "Poste analyste", "client MCP, jeton, posture", GOLD),
        ("ZONE 1", "Serveur non privilégié", "CP1 à CP7, un seul journal", BLUE),
        ("ZONE 2", "Données classifiées", "SQLite 0440, clé 0400, audit", GREEN),
        ("ZONE 3", "Flux partenaire", "non fiable, neutralisé à l'ingestion", RED),
    ]
    for index, (zid, label, detail, accent) in enumerate(zones):
        yy = CONTENT_Y + index * 1.03
        card(slide, 8.75, yy, 3.9, 0.9, LIGHT)
        rect(slide, 8.75, yy, 0.06, 0.9, accent)
        text(slide, 8.95, yy + 0.08, 0.7, 0.25, zid, size=8.5, bold=True, color=accent)
        text(slide, 9.68, yy + 0.07, 2.8, 0.3, label, size=11.5, bold=True, color=NAVY)
        text(slide, 9.68, yy + 0.39, 2.75, 0.4, detail, size=8.5, color=MUTED)
    card(slide, 8.75, 6.1, 3.9, 0.55, PALE_GOLD, border=GOLD)
    text(slide, 8.94, 6.2, 3.55, 0.35, "Fail-closed : pas d'audit = pas de réponse.",
         size=9.4, bold=True, color=NAVY)
    notes(slide, s)


def slide_09(prs, s):
    slide = chrome(prs, s)
    add_image(slide, FIG / "fig_parcours_document.png", 0.42, CONTENT_Y, 5.95, 4.72)
    text(slide, 0.55, 6.75, 5.7, 0.22,
         "Source : conception du projet — chaque étape cite une exigence.",
         size=7.8, italic=True, color=MUTED)
    steps = [
        ("1–2", "Réception et validation", "contenu non exécutable", GOLD),
        ("3–4", "Étiquetage et stockage", "niveau, TLP, compartiments", BLUE),
        ("5–6", "Demande et décision", "jeton, posture, ABAC", GREEN),
        ("7–8", "Enveloppe et réponse", "masquage, fanions, audit", RED),
    ]
    for index, (num, label, detail, accent) in enumerate(steps):
        x = 6.65 + (index % 2) * 2.95
        yy = CONTENT_Y + (index // 2) * 1.55
        card(slide, x, yy, 2.75, 1.35, LIGHT)
        text(slide, x + 0.16, yy + 0.12, 0.7, 0.3, num, size=9, bold=True, color=accent)
        text(slide, x + 0.16, yy + 0.4, 2.4, 0.5, label, size=10.5, bold=True, color=NAVY, line_spacing=1.0)
        text(slide, x + 0.16, yy + 0.95, 2.35, 0.3, detail, size=8.2, color=MUTED)
    card(slide, 6.65, 5.2, 5.75, 0.92, PALE_RED, border=RED)
    text(slide, 6.86, 5.31, 5.3, 0.7,
         "Refus : message générique pour le client, motif précis dans l'audit. Aucun oracle.",
         size=10.2, bold=True, color=NAVY, line_spacing=1.03)
    card(slide, 6.65, 6.25, 5.75, 0.4, NAVY)
    text(slide, 6.82, 6.31, 5.4, 0.25,
         "Source visible au niveau 3 seulement ; canari absent de tout refus.",
         size=8.2, color=WHITE)
    notes(slide, s)


def slide_10(prs, s):
    slide = chrome(prs, s)
    card(slide, M, CONTENT_Y, 6.0, 2.45, LIGHT)
    text(slide, M + 0.2, CONTENT_Y + 0.13, 5.5, 0.28, "SOCLE — 9 MODULES, BIBLIOTHÈQUE STANDARD", size=8.6, bold=True, color=BLUE)
    modules = ["config", "labels", "sanitize", "auth", "policy", "store", "intel", "tools", "server"]
    for index, module in enumerate(modules):
        x = M + 0.2 + (index % 3) * 1.85
        yy = CONTENT_Y + 0.55 + (index // 3) * 0.52
        chip(slide, x, yy, 1.65, 0.38, module, PALE_BLUE, NAVY, size=8.6)
    card(slide, 6.9, CONTENT_Y, 5.73, 2.45, LIGHT)
    text(slide, 7.1, CONTENT_Y + 0.13, 5.3, 0.28, "CONTRAT — 6 OUTILS EN LECTURE SEULE", size=8.6, bold=True, color=GREEN)
    tools = ["ti_search", "ti_get_advisory", "ti_enrich_ioc", "ti_correlate", "ti_score_asset", "ti_health"]
    for index, tool in enumerate(tools):
        x = 7.1 + (index % 2) * 2.6
        yy = CONTENT_Y + 0.55 + (index // 2) * 0.54
        chip(slide, x, yy, 2.4, 0.4, tool, PALE_GREEN, GREEN, size=8.8)
    text(slide, 7.1, CONTENT_Y + 2.18, 5.2, 0.25,
         "schémas stricts · additionalProperties:false · readOnlyHint", size=7.8, color=MUTED)
    rect(slide, M, 4.75, 11.95, 0.018, LINE)
    steps = [
        ("12:40", "socle"), ("12:52", "données"), ("13:02", "serveur"),
        ("13:08", "échec réel"), ("13:14", "correction"), ("13:16", "campagne verte"),
    ]
    for index, (hour, label) in enumerate(steps):
        x = M + 0.35 + index * 1.93
        rect(slide, x, 5.28, 0.035, 0.7, GOLD)
        text(slide, x - 0.15, 5.03, 0.9, 0.22, hour, size=8.8, bold=True, color=GOLD)
        text(slide, x - 0.02, 5.9, 1.65, 0.35, label, size=9.2, color=INK)
        if index < len(steps) - 1:
            rect(slide, x + 0.05, 5.59, 1.88, 0.018, LINE)
    card(slide, M, 6.34, 11.95, 0.5, PALE_GOLD, border=GOLD)
    text(slide, M + 0.2, 6.42, 11.5, 0.32,
         "Premier échec à 13 h 08 : décodage UTF-8 de l'audit ; corrigé à 13 h 14, rejoué, conservé.",
         size=9.2, bold=True, color=NAVY)
    notes(slide, s)


def slide_11(prs, s):
    slide = chrome(prs, s)
    add_image(slide, CAP_ECRAN / "07_moindre_privilege.png", 0.45, CONTENT_Y + 0.35, 8.7, 2.4)
    text(slide, 0.55, 4.77, 8.5, 0.22,
         "Capture d'écran réelle : exécution sous « nobody », écriture refusée, aucun socket.",
         size=7.8, italic=True, color=MUTED)
    kpis = [("65534", "uid du service", "nobody", GREEN), ("0440", "base SQLite", "lecture seule", BLUE),
            ("0400", "clé HMAC", "propriétaire", GOLD), ("0", "port réseau", "stdio local", RED)]
    for index, (value, label, detail, accent) in enumerate(kpis):
        x = 9.45 + (index % 2) * 1.65
        yy = CONTENT_Y + (index // 2) * 1.32
        card(slide, x, yy, 1.52, 1.14, LIGHT)
        text(slide, x + 0.12, yy + 0.1, 1.35, 0.4, value, size=19, bold=True, color=accent)
        text(slide, x + 0.12, yy + 0.52, 1.3, 0.22, label, size=7.6, bold=True, color=INK)
        text(slide, x + 0.12, yy + 0.78, 1.3, 0.22, detail, size=7.2, color=MUTED)
    status = [("TESTÉ", "uid 65534 · 0440 · 0400", GREEN),
              ("PROPOSÉ", "unité systemd", GOLD),
              ("NON TESTÉ", "HTTP · mTLS · EDR réel", RED)]
    for index, (label, detail, color) in enumerate(status):
        x = M + index * 4.0
        card(slide, x, 5.2, 3.75, 0.86, LIGHT)
        chip(slide, x + 0.15, 5.3, 1.03, 0.32, label, color, WHITE, size=7.4)
        text(slide, x + 1.3, 5.32, 2.3, 0.55, detail, size=8.8, color=INK)
    card(slide, M, 6.25, 11.95, 0.5, PALE_GOLD, border=GOLD)
    text(slide, M + 0.2, 6.34, 11.5, 0.32,
         "Le service ne peut pas écrire la base ; l'ingestion est une opération séparée.",
         size=9.8, bold=True, color=NAVY)
    notes(slide, s)


def slide_12(prs, s):
    slide = chrome(prs, s)
    add_image(slide, CAP_ECRAN / "03_fuites_bloquees.png", 0.48, CONTENT_Y, 11.9, 3.95)
    text(slide, 0.55, 6.05, 11.6, 0.22,
         "Source : preuves/logs/demo_capture.json — scènes S2 et S3 ; canari absent des deux réponses.",
         size=7.8, italic=True, color=MUTED)
    chips = [("Confidentiel → Secret : refus", RED), ("Secret hors compartiment : refus", RED),
             ("Deux motifs d'audit distincts", GREEN)]
    for index, (label, color) in enumerate(chips):
        chip(slide, M + index * 4.0, 6.42, 3.75, 0.4, label, color, WHITE, size=8.4)
    notes(slide, s)


def slide_13(prs, s):
    slide = chrome(prs, s)
    add_image(slide, CAP_ECRAN / "04_acces_et_injection.png", 0.5, 1.9, 6.6, 2.21)
    add_image(slide, CAP_ECRAN / "05_audit_falsification.png", 0.5, 4.25, 5.2, 2.0)
    text(slide, 0.55, 6.3, 8.25, 0.22,
         "Source : démonstration S5 et S6 — instruction neutralisée, chaîne vérifiée.",
         size=7.6, italic=True, color=MUTED)
    card(slide, 9.15, CONTENT_Y, 3.48, 1.35, PALE_RED, border=RED)
    text(slide, 9.35, CONTENT_Y + 0.12, 2.9, 0.25, "INJECTION", size=8.5, bold=True, color=RED)
    text(slide, 9.35, CONTENT_Y + 0.4, 1.4, 0.6, "3 / 3", size=29, bold=True, color=RED)
    text(slide, 10.6, CONTENT_Y + 0.58, 1.8, 0.3, "charges neutralisées", size=8.2, color=INK)
    text(slide, 9.35, CONTENT_Y + 1.02, 2.9, 0.25, "0 faux positif sur 11 bénins", size=8.4, bold=True, color=GREEN)
    card(slide, 9.15, 3.5, 3.48, 1.4, PALE_GREEN, border=GREEN)
    text(slide, 9.35, 3.63, 2.9, 0.25, "AUDIT", size=8.5, bold=True, color=GREEN)
    text(slide, 9.35, 3.91, 1.45, 0.6, "116", size=29, bold=True, color=GREEN)
    text(slide, 10.55, 4.09, 1.9, 0.3, "entrées chaînées", size=8.2, color=INK)
    text(slide, 9.35, 4.55, 2.9, 0.25, "115 appels + 1 erreur protocole", size=8.3, bold=True, color=INK)
    card(slide, 9.15, 5.15, 3.48, 1.1, PALE_GOLD, border=GOLD)
    text(slide, 9.33, 5.27, 3.1, 0.8,
         "La chaîne SHA-256 est vérifiée en une commande ; la source est masquée.",
         size=9.2, bold=True, color=NAVY, line_spacing=1.05)
    notes(slide, s)


def slide_14(prs, s):
    slide = chrome(prs, s)
    add_image(slide, CAP / "cap_07_campagne_tests.png", 0.45, CONTENT_Y, 8.2, 4.08)
    text(slide, 0.55, 6.15, 8.0, 0.22,
         "Source : preuves/logs/tests_latest.json — campagne du 27 octobre 2026.",
         size=7.8, italic=True, color=MUTED)
    kpis = [("31 / 31", "scénarios réussis", GREEN), ("16 / 16", "abus maîtrisés", GREEN),
            ("12,42 ms", "latence p95", BLUE), ("10,0 ms", "révocation", GOLD)]
    for index, (value, label, color) in enumerate(kpis):
        yy = CONTENT_Y + index * 1.06
        card(slide, 9.0, yy, 3.63, 0.9, LIGHT)
        rect(slide, 9.0, yy, 0.06, 0.9, color)
        text(slide, 9.2, yy + 0.11, 1.9, 0.4, value, size=21, bold=True, color=color)
        text(slide, 11.05, yy + 0.25, 1.45, 0.3, label, size=8.4, color=MUTED, align=PP_ALIGN.RIGHT)
    card(slide, M, 6.42, 11.95, 0.48, PALE_GREEN, border=GREEN)
    text(slide, M + 0.2, 6.51, 11.5, 0.3,
         "Mêmes données, mêmes outils : la maquette naïve échoue 8 / 8, le serveur durci échoue 0 / 8.",
         size=9.5, bold=True, color=NAVY)
    notes(slide, s)


def slide_15(prs, s):
    slide = chrome(prs, s)
    card(slide, M, CONTENT_Y, 5.35, 3.95, PALE_GREEN, border=GREEN)
    text(slide, M + 0.28, CONTENT_Y + 0.18, 4.8, 0.28, "CONTRÔLES COMPARÉS", size=8.8, bold=True, color=GREEN)
    text(slide, M + 0.25, CONTENT_Y + 0.55, 2.0, 1.0, "8 → 0", size=53, bold=True, color=GREEN)
    text(slide, M + 2.55, CONTENT_Y + 0.92, 2.4, 0.55, "faiblesses\nsur 8", size=13, bold=True, color=NAVY)
    text(slide, M + 0.28, CONTENT_Y + 1.85, 4.8, 0.55,
         "Canari, source, injection, validation, séparation, journal, révocation, débit.",
         size=9.8, color=INK, line_spacing=1.1)
    text(slide, M + 0.28, CONTENT_Y + 2.65, 4.8, 0.8,
         "La maquette naïve expose tout ; le serveur durci refuse, masque ou neutralise.",
         size=9.5, italic=True, color=MUTED, line_spacing=1.08)
    card(slide, 6.45, CONTENT_Y, 6.2, 3.95, LIGHT)
    text(slide, 6.68, CONTENT_Y + 0.18, 5.7, 0.28, "CE QUI N'EST PAS TESTÉ", size=8.8, bold=True, color=RED)
    limits = [
        ("Isolation réseau", "proposée", GOLD),
        ("Agent EDR réel", "déclaratif", GOLD),
        ("Ancrage externe audit", "absent", RED),
        ("HTTP · OAuth 2.1 · mTLS", "non testé", RED),
        ("Charge réelle", "non testée", GOLD),
    ]
    yy = CONTENT_Y + 0.6
    for label, status, color in limits:
        text(slide, 6.68, yy, 3.7, 0.25, label, size=10.2, bold=True, color=INK)
        chip(slide, 10.35, yy - 0.04, 1.35, 0.29, status.upper(), color, WHITE, size=6.9)
        yy += 0.6
    card(slide, M, 6.12, 11.95, 0.55, PALE_GOLD, border=GOLD)
    text(slide, M + 0.2, 6.22, 11.5, 0.34,
         "Une limite nommée n'est pas un échec caché : c'est un travail futur priorisé.",
         size=10.2, bold=True, color=NAVY)
    notes(slide, s)


def slide_16(prs, s):
    slide = chrome(prs, s, dark=True)
    text(slide, M, 1.95, 11.9, 0.35, "LA THÈSE, EN UNE PHRASE", size=9.5, bold=True, color=GOLD)
    text(slide, M, 2.38, 11.8, 1.35,
         "Un serveur MCP peut rester simple, local et honnête\nen rendant la protection du secret démontrable.",
         size=26, bold=True, color=WHITE, line_spacing=1.08)
    pillars = [
        ("DÉCIDER", "ABAC, habilitation,\ncompartiments", GOLD),
        ("ÉTIQUETER", "niveau, TLP, masquage,\nfanions d'injection", BLUE),
        ("TRACER", "journal chaîné,\nrevue et révocation", GREEN),
    ]
    for index, (title, body, accent) in enumerate(pillars):
        x = M + index * 4.0
        card(slide, x, 4.12, 3.75, 1.55, RGBColor(0x12, 0x2A, 0x4A), border=accent)
        text(slide, x + 0.2, 4.28, 3.3, 0.3, title, size=9.5, bold=True, color=accent)
        text(slide, x + 0.2, 4.62, 3.3, 0.8, body, size=10.2, color=WHITE, line_spacing=1.08)
    text(slide, M, 6.0, 5.8, 0.35, "31 / 31 · 16 / 16 · 12,42 ms · 8 → 0", size=15, bold=True, color=GOLD)
    text(slide, 8.0, 5.98, 4.6, 0.35, "Merci — questions bienvenues.", size=12, color=WHITE, align=PP_ALIGN.RIGHT)
    rect(slide, M, 6.55, 11.95, 0.015, RGBColor(0x3B, 0x52, 0x70))
    text(slide, M, 6.7, 11.9, 0.28,
         "Rapport, preuves, captures et commandes de reproduction : dossier du projet.",
         size=8.2, color=RGBColor(0xAF, 0xBE, 0xD0))
    notes(slide, s)


BUILDERS = {
    1: title_slide, 2: slide_02, 3: slide_03, 4: slide_04, 5: slide_05,
    6: slide_06, 7: slide_07, 8: slide_08, 9: slide_09, 10: slide_10,
    11: slide_11, 12: slide_12, 13: slide_13, 14: slide_14, 15: slide_15,
    16: slide_16,
}


def main() -> int:
    prs = Presentation()
    prs.slide_width = Inches(W)
    prs.slide_height = Inches(H)
    for s in SLIDES:
        BUILDERS[s["n"]](prs, s)
    prs.core_properties.title = "Serveur MCP intelligent de renseignement sur les menaces"
    prs.core_properties.author = "Mzabi Sirine · Khaled Ellif"
    prs.core_properties.subject = "Soutenance TEK-UP SSIR — Art de protection du secret & Cyber OPS"
    prs.core_properties.comments = "Deck 16:9, 16 slides, notes orales et questions jury."
    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUT)
    print(f"DECK: {OUT}")
    print(f"SLIDES: {len(prs.slides)} · NOTES: {sum(1 for _ in prs.slides)}")
    subprocess.run(
        ["libreoffice", "--headless", "--convert-to", "pdf", "--outdir", str(OUT.parent), str(OUT)],
        check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL,
    )
    print(f"PDF : {OUT.with_suffix('.pdf')}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
