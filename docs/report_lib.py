"""Outillage de mise en page du rapport (python-docx, champs Word réels)."""
from __future__ import annotations

import re
from pathlib import Path
from typing import Iterable, Sequence

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image

NAVY = RGBColor(0x10, 0x24, 0x3F)
BLUE = RGBColor(0x2F, 0x6D, 0xB3)
GREEN = RGBColor(0x2C, 0x7A, 0x5B)
AMBER = RGBColor(0xB7, 0x79, 0x1F)
RED = RGBColor(0xA8, 0x32, 0x32)
GREY = RGBColor(0x5B, 0x65, 0x77)
INK = RGBColor(0x1B, 0x24, 0x33)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

NAVY_HEX = "10243F"
BLUE_HEX = "2F6DB3"
LIGHT_HEX = "EEF4FB"
ZEBRA_HEX = "F3F6FA"
GREEN_HEX = "E7F2EC"
AMBER_HEX = "FBF1DC"
RED_HEX = "FBE9E7"
GREY_HEX = "F4F6F9"


def _set_run_font(run, name: str, size: float | None = None, bold: bool | None = None,
                  italic: bool | None = None, color: RGBColor | None = None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = color


def _field_runs(paragraph, instruction: str, cached: str, *, bold: bool = False,
                size: float | None = None, color: RGBColor | None = None):
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")

    run_begin = paragraph.add_run()
    run_begin._r.append(begin)
    run_instr = paragraph.add_run()
    run_instr._r.append(instr)
    run_sep = paragraph.add_run()
    run_sep._r.append(separate)
    run_value = paragraph.add_run(cached)
    run_end = paragraph.add_run()
    run_end._r.append(end)
    if size:
        for run in (run_value,):
            run.font.size = Pt(size)
    if bold:
        run_value.bold = True
    if color:
        run_value.font.color.rgb = color


def _shade(paragraph_or_cell, fill: str):
    element = paragraph_or_cell._p if hasattr(paragraph_or_cell, "_p") else paragraph_or_cell._tc
    pr = element.get_or_add_pPr() if hasattr(element, "get_or_add_pPr") else element.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    pr.append(shd)


def _border(paragraph, *, top: bool = False, bottom: bool = False, left: bool = False,
            color: str = NAVY_HEX, size: int = 8):
    pPr = paragraph._p.get_or_add_pPr()
    pBdr = pPr.find(qn("w:pBdr"))
    if pBdr is None:
        pBdr = OxmlElement("w:pBdr")
        pPr.insert_element_before(
            pBdr, "w:shd", "w:tabs", "w:suppressAutoHyphens", "w:kinsoku", "w:wordWrap",
            "w:overflowPunct", "w:topLinePunct", "w:autoSpaceDE", "w:autoSpaceDN", "w:bidi",
            "w:adjustRightInd", "w:snapToGrid", "w:spacing", "w:ind", "w:contextualSpacing",
            "w:mirrorIndents", "w:suppressOverlap", "w:jc", "w:textDirection",
            "w:textAlignment", "w:textboxTightWrap", "w:outlineLvl", "w:divId",
            "w:cnfStyle", "w:rPr", "w:sectPr", "w:pPrChange",
        )
    for side, active in (("top", top), ("bottom", bottom), ("left", left)):
        if not active:
            continue
        element = OxmlElement(f"w:{side}")
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), str(size))
        element.set(qn("w:space"), "4")
        element.set(qn("w:color"), color)
        pBdr.append(element)


def _cell_shade(cell, fill: str):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def _cell_margins(table, top: int = 40, bottom: int = 40, left: int = 80, right: int = 80):
    tblPr = table._tbl.tblPr
    mar = OxmlElement("w:tblCellMar")
    for side, value in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        element = OxmlElement(f"w:{side}")
        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")
        mar.append(element)
    tblPr.append(mar)


def _repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    trPr.append(header)
    cant = OxmlElement("w:cantSplit")
    cant.set(qn("w:val"), "true")
    trPr.append(cant)


def _no_split(row):
    trPr = row._tr.get_or_add_trPr()
    cant = OxmlElement("w:cantSplit")
    cant.set(qn("w:val"), "true")
    trPr.append(cant)


class Report:
    def __init__(self, emit: bool = True):
        self.emit = emit
        self.doc = Document()
        self.headings: list[tuple[int, str]] = []
        self.figures: list[str] = []
        self.tables: list[str] = []
        self.figure_counter = 0
        self.table_counter = 0
        self._setup()

    # -- mise en page ------------------------------------------------------
    def _setup(self) -> None:
        section = self.doc.sections[0]
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)
        section.left_margin = Cm(2.0)
        section.right_margin = Cm(2.0)
        section.top_margin = Cm(2.0)
        section.bottom_margin = Cm(1.8)
        section.header_distance = Cm(0.9)
        section.footer_distance = Cm(0.8)
        section.different_first_page_header_footer = True

        styles = self.doc.styles
        normal = styles["Normal"]
        normal.font.name = "Times New Roman"
        normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        normal.font.size = Pt(11.2)
        normal.font.color.rgb = INK
        normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        normal.paragraph_format.line_spacing = 1.16
        normal.paragraph_format.space_after = Pt(5)
        normal.paragraph_format.widow_control = True

        for name, size, color, before, after in (
            ("Heading 1", 17, NAVY, 18, 10),
            ("Heading 2", 14, BLUE, 14, 6),
            ("Heading 3", 12, NAVY, 10, 4),
        ):
            style = styles[name]
            style.font.name = "Calibri"
            style._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
            style.font.size = Pt(size)
            style.font.bold = True
            style.font.color.rgb = color
            style.paragraph_format.space_before = Pt(before)
            style.paragraph_format.space_after = Pt(after)
            style.paragraph_format.keep_with_next = True
            style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
            style.paragraph_format.line_spacing = 1.05

        caption = styles.add_style("Légende", 1)
        caption.base_style = styles["Normal"]
        caption.font.name = "Calibri"
        caption.font.size = Pt(9)
        caption.font.italic = True
        caption.font.color.rgb = GREY
        caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        caption.paragraph_format.space_before = Pt(2)
        caption.paragraph_format.space_after = Pt(2)

        source = styles.add_style("Source", 1)
        source.base_style = styles["Normal"]
        source.font.name = "Calibri"
        source.font.size = Pt(8)
        source.font.italic = True
        source.font.color.rgb = GREY
        source.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        source.paragraph_format.space_after = Pt(10)

        code = styles.add_style("Code", 1)
        code.base_style = styles["Normal"]
        code.font.name = "Consolas"
        code._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
        code.font.size = Pt(8.6)
        code.font.color.rgb = RGBColor(0x1F, 0x2A, 0x3A)
        code.paragraph_format.left_indent = Cm(0.4)
        code.paragraph_format.right_indent = Cm(0.2)
        code.paragraph_format.space_before = Pt(3)
        code.paragraph_format.space_after = Pt(6)
        code.paragraph_format.line_spacing = 1.05
        code.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT

        front = styles.add_style("Titre liminaire", 1)
        front.base_style = styles["Normal"]
        front.font.name = "Calibri"
        front.font.size = Pt(20)
        front.font.bold = True
        front.font.color.rgb = NAVY
        front.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        front.paragraph_format.space_before = Pt(4)
        front.paragraph_format.space_after = Pt(12)

        for level, size in ((1, 10.5), (2, 10.0), (3, 9.5)):
            style = styles.add_style(f"TOC {level}", 1)
            style.base_style = styles["Normal"]
            style.font.name = "Calibri"
            style.font.size = Pt(size)
            style.font.color.rgb = INK
            style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
            style.paragraph_format.space_after = Pt(2)
            style.paragraph_format.line_spacing = 1.0

        self._header_footer(section)

        props = self.doc.core_properties
        props.title = "Serveur MCP intelligent de Threat Intelligence — protection du secret"
        props.author = "Mzabi Sirine · Khaled Ellif"
        props.subject = "Projet TEK-UP SSIR — Art de protection du secret & Cyber OPS"
        props.keywords = "MCP, Threat Intelligence, ABAC, étiquetage, audit, prompt injection"
        props.comments = "Rapport généré par python-docx avec champs Word réels."

        settings = self.doc.settings.element
        update = OxmlElement("w:updateFields")
        update.set(qn("w:val"), "true")
        settings.append(update)

    def _header_footer(self, section) -> None:
        header = section.header
        p = header.paragraphs[0]
        p.paragraph_format.tab_stops.add_tab_stop(Cm(16.8), WD_TAB_ALIGNMENT.RIGHT)
        run = p.add_run("TEK-UP · SSIR 2026–2027")
        _set_run_font(run, "Calibri", 8.5, color=GREY)
        run = p.add_run("\tServeur MCP de Threat Intelligence")
        _set_run_font(run, "Calibri", 8.5, color=GREY)
        _border(p, bottom=True, color="C9D1DE", size=4)

        footer = section.footer
        p = footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run("Page ")
        _set_run_font(run, "Calibri", 8.5, color=GREY)
        _field_runs(p, " PAGE ", "1", size=8.5, color=GREY)
        run = p.add_run(" sur ")
        _set_run_font(run, "Calibri", 8.5, color=GREY)
        _field_runs(p, " NUMPAGES ", "1", size=8.5, color=GREY)

        # La page de couverture n'affiche ni en-tête ni pied de page.
        first_header = section.first_page_header
        first_header.paragraphs[0].text = ""
        first_footer = section.first_page_footer
        first_footer.paragraphs[0].text = ""

    # -- texte -------------------------------------------------------------
    def _rich(self, paragraph, text: str, size: float | None = None, color: RGBColor | None = None):
        bold = False
        italic = False
        for token in re.split(r"(`[^`]+`|\*\*|__)", text):
            if token == "**":
                bold = not bold
                continue
            if token == "__":
                italic = not italic
                continue
            if not token:
                continue
            if token.startswith("`") and token.endswith("`"):
                run = paragraph.add_run(token[1:-1])
                _set_run_font(run, "Consolas", (size - 0.6) if size else None,
                              bold=bold, italic=italic,
                              color=color or RGBColor(0x1F, 0x2A, 0x3A))
                continue
            run = paragraph.add_run(token)
            run.bold = bold
            run.italic = italic
            if size:
                run.font.size = Pt(size)
            if color:
                run.font.color.rgb = color

    def p(self, text: str, *, size: float | None = None, italic: bool = False,
          align: WD_ALIGN_PARAGRAPH | None = None, space_after: float | None = None,
          color: RGBColor | None = None, keep_with_next: bool = False):
        paragraph = self.doc.add_paragraph()
        if align is not None:
            paragraph.alignment = align
        if space_after is not None:
            paragraph.paragraph_format.space_after = Pt(space_after)
        if italic:
            for run in re.split(r"(\*\*|__)", text):
                if run in ("**", "__") or not run:
                    continue
                r = paragraph.add_run(run)
                _set_run_font(r, "Times New Roman", size or 11.5, italic=True, color=color)
        else:
            self._rich(paragraph, text, size=size, color=color)
        paragraph.paragraph_format.keep_with_next = keep_with_next
        return paragraph

    def lead(self, text: str):
        """Paragraphe d'ouverture légèrement plus grand."""
        return self.p(text, size=12.2)

    def quote(self, text: str, author: str | None = None):
        p = self.doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(1.0)
        p.paragraph_format.right_indent = Cm(1.0)
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(4)
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run = p.add_run("« " + text + " »")
        _set_run_font(run, "Times New Roman", 11.5, italic=True, color=NAVY)
        if author:
            a = self.doc.add_paragraph()
            a.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            a.paragraph_format.right_indent = Cm(1.0)
            a.paragraph_format.space_after = Pt(10)
            run = a.add_run("— " + author)
            _set_run_font(run, "Calibri", 9, color=GREY)

    def code(self, text: str):
        for line in _split_code(text):
            p = self.doc.add_paragraph(style="Code")
            _shade(p, GREY_HEX)
            _border(p, left=True, color="2F6DB3", size=12)
            run = p.add_run(line if line else " ")
            _set_run_font(run, "Consolas", 8.6, color=RGBColor(0x1F, 0x2A, 0x3A))

    def callout(self, title: str, text: str, *, fill: str = LIGHT_HEX,
                border: str = BLUE_HEX):
        table = self.doc.add_table(rows=1, cols=1)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        _cell_margins(table, 90, 90, 140, 140)
        cell = table.cell(0, 0)
        cell.width = Cm(16.6)
        _cell_shade(cell, fill)
        tcPr = cell._tc.get_or_add_tcPr()
        borders = OxmlElement("w:tcBorders")
        for side in ("top", "bottom", "right"):
            element = OxmlElement(f"w:{side}")
            element.set(qn("w:val"), "nil")
            borders.append(element)
        left = OxmlElement("w:left")
        left.set(qn("w:val"), "single")
        left.set(qn("w:sz"), "18")
        left.set(qn("w:color"), border)
        borders.append(left)
        tcPr.append(borders)
        paragraph = cell.paragraphs[0]
        paragraph.paragraph_format.space_after = Pt(2)
        run = paragraph.add_run(title)
        _set_run_font(run, "Calibri", 9.5, bold=True, color=NAVY)
        second = cell.add_paragraph()
        second.paragraph_format.space_after = Pt(0)
        self._rich(second, text, size=9.6)

    def callout_ok(self, title: str, text: str):
        self.callout(title, text, fill=GREEN_HEX, border="2C7A5B")

    def callout_warn(self, title: str, text: str):
        self.callout(title, text, fill=RED_HEX, border="A83232")

    # -- titres ------------------------------------------------------------
    def part(self, title: str, subtitle: str | None = None):
        p = self.doc.add_paragraph(style="Heading 1")
        p.paragraph_format.page_break_before = True
        p.paragraph_format.space_before = Pt(90)
        p.paragraph_format.space_after = Pt(8)
        run = p.add_run(title)
        _set_run_font(run, "Calibri", 24, bold=True, color=NAVY)
        self.headings.append((1, title))
        _border(p, bottom=True, color="B08A3E", size=12)
        if subtitle:
            s = self.doc.add_paragraph()
            s.paragraph_format.space_before = Pt(8)
            s.paragraph_format.space_after = Pt(24)
            run = s.add_run(subtitle)
            _set_run_font(run, "Calibri", 11.5, italic=True, color=GREY)

    def h1(self, text: str, page_break: bool = False):
        p = self.doc.add_paragraph(style="Heading 1")
        if page_break:
            p.paragraph_format.page_break_before = True
        run = p.add_run(text)
        _set_run_font(run, "Calibri", 17, bold=True, color=NAVY)
        self.headings.append((1, text))
        return p

    def h2(self, text: str):
        p = self.doc.add_paragraph(style="Heading 2")
        run = p.add_run(text)
        _set_run_font(run, "Calibri", 14, bold=True, color=BLUE)
        self.headings.append((2, text))
        return p

    def h3(self, text: str):
        p = self.doc.add_paragraph(style="Heading 3")
        run = p.add_run(text)
        _set_run_font(run, "Calibri", 12, bold=True, color=NAVY)
        self.headings.append((3, text))
        return p

    def front_title(self, text: str, page_break: bool = True):
        if page_break:
            self.doc.add_page_break()
        p = self.doc.add_paragraph(style="Titre liminaire")
        run = p.add_run(text)
        _set_run_font(run, "Calibri", 20, bold=True, color=NAVY)
        _border(p, bottom=True, color="B08A3E", size=8)
        return p

    # -- légendes, figures, tableaux ---------------------------------------
    def _caption(self, kind: str, text: str) -> str:
        if kind == "Figure":
            self.figure_counter += 1
            number = self.figure_counter
            self.figures.append(f"Figure {number} — {text}")
        else:
            self.table_counter += 1
            number = self.table_counter
            self.tables.append(f"Tableau {number} — {text}")
        p = self.doc.add_paragraph(style="Légende")
        p.paragraph_format.keep_with_next = kind in ("Tableau", "Figure")
        run = p.add_run(f"{kind} ")
        _set_run_font(run, "Calibri", 9, italic=True, color=GREY)
        _field_runs(p, f" SEQ {kind} \\* ARABIC ", str(number), bold=False,
                    size=9, color=GREY)
        run = p.add_run(" — " + text)
        _set_run_font(run, "Calibri", 9, italic=True, color=GREY)
        return f"{kind} {number} — {text}"

    def source(self, text: str):
        p = self.doc.add_paragraph(style="Source")
        for token in re.split(r"(`[^`]+`)", text):
            if not token:
                continue
            if token.startswith("`") and token.endswith("`"):
                run = p.add_run(token[1:-1])
                _set_run_font(run, "Consolas", 7.6, italic=True, color=GREY)
            else:
                run = p.add_run(token)
                _set_run_font(run, "Calibri", 8, italic=True, color=GREY)

    def fig_band(self, tag: str, title: str, note: str = ""):
        p = self.doc.add_paragraph()
        pf = p.paragraph_format
        pf.space_before = Pt(10)
        pf.space_after = Pt(3)
        pf.keep_with_next = True
        pf.tab_stops.add_tab_stop(Cm(16.4), WD_TAB_ALIGNMENT.RIGHT)
        run = p.add_run(tag)
        _set_run_font(run, "Calibri", 7.8, bold=True, color=BLUE)
        spacing = OxmlElement("w:spacing")
        spacing.set(qn("w:val"), "30")
        run._element.get_or_add_rPr().insert_element_before(
            spacing, "w:w", "w:kern", "w:position", "w:sz", "w:szCs", "w:highlight", "w:u",
            "w:effect", "w:bdr", "w:shd", "w:fitText", "w:vertAlign", "w:rtl", "w:cs",
            "w:em", "w:lang", "w:eastAsianLayout", "w:specVanish", "w:oMath",
        )
        run = p.add_run("   " + title)
        _set_run_font(run, "Calibri", 9, bold=True, color=NAVY)
        if note:
            run = p.add_run("\t" + note)
            _set_run_font(run, "Calibri", 8, italic=True, color=GREY)
        _border(p, bottom=True, color="C9CED6", size=4)
        return p

    def fig(self, image_path: str | Path, caption: str, source: str,
            width_cm: float | None = None):
        path = Path(image_path)
        with Image.open(path) as img:
            width_px, height_px = img.size
        if width_cm is None:
            ratio = height_px / max(width_px, 1)
            width_cm = 13.2 if ratio > 0.95 else 16.2
            if ratio > 1.35:
                width_cm = 11.8
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.space_before = Pt(5)
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run()
        shape = run.add_picture(str(path), width=Cm(width_cm))
        outline = OxmlElement("a:ln")
        outline.set("w", "9525")
        solid = OxmlElement("a:solidFill")
        color = OxmlElement("a:srgbClr")
        color.set("val", "C9CED6")
        solid.append(color)
        outline.append(solid)
        shape._inline.find(".//" + qn("pic:spPr")).append(outline)
        self._caption("Figure", caption)
        self.source(source)

    def table(self, headers: Sequence[str], rows: Iterable[Sequence[str]], caption: str | None,
              widths: Sequence[float], source: str | None = None, font_size: float = 8.3,
              header_size: float = 8.3, caption_above: bool = True, zebra: bool = True,
              first_col_bold: bool = False, align_center_cols: Sequence[int] = ()):
        if caption:
            if caption_above:
                self._caption("Tableau", caption)
            else:
                after_caption = caption
        table = self.doc.add_table(rows=1, cols=len(headers))
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        _cell_margins(table)
        header = table.rows[0]
        _repeat_header(header)
        for index, (text, width) in enumerate(zip(headers, widths)):
            cell = header.cells[index]
            cell.width = Cm(width)
            _cell_shade(cell, NAVY_HEX)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = paragraph.add_run(text)
            _set_run_font(run, "Calibri", header_size, bold=True, color=WHITE)

        for row_index, row in enumerate(rows):
            cells = table.add_row().cells
            _no_split(table.rows[-1])
            for index, (text, width) in enumerate(zip(row, widths)):
                cell = cells[index]
                cell.width = Cm(width)
                if zebra and row_index % 2 == 1:
                    _cell_shade(cell, ZEBRA_HEX)
                paragraph = cell.paragraphs[0]
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.02
                paragraph.alignment = (WD_ALIGN_PARAGRAPH.CENTER if index in align_center_cols
                                       else WD_ALIGN_PARAGRAPH.LEFT)
                self._rich(paragraph, str(text), size=font_size)
                if first_col_bold and index == 0:
                    for run in paragraph.runs:
                        run.bold = True

        if caption and not caption_above:
            self._caption("Tableau", after_caption)
        if source:
            self.source(source)
        else:
            self.doc.add_paragraph().paragraph_format.space_after = Pt(2)

    # -- sommaire et listes ------------------------------------------------
    def toc(self, entries: Sequence[tuple[int, str]], pages: dict[str, str],
            instruction: str, *, entry_kind: str = "heading"):
        p = self.doc.add_paragraph()
        begin = OxmlElement("w:fldChar")
        begin.set(qn("w:fldCharType"), "begin")
        instr = OxmlElement("w:instrText")
        instr.set(qn("xml:space"), "preserve")
        instr.text = instruction
        separate = OxmlElement("w:fldChar")
        separate.set(qn("w:fldCharType"), "separate")
        for element in (begin, instr, separate):
            run = p.add_run()
            run._r.append(element)
        p.paragraph_format.space_after = Pt(0)

        for index, (level, text) in enumerate(entries):
            paragraph = self.doc.add_paragraph(style=f"TOC {min(level, 3)}")
            paragraph.paragraph_format.left_indent = Cm(0.0 if level == 1 else (0.55 if level == 2 else 1.1))
            paragraph.paragraph_format.tab_stops.add_tab_stop(
                Cm(16.4), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
            if level == 1:
                paragraph.paragraph_format.space_before = Pt(5 if index else 0)
            run = paragraph.add_run(text)
            _set_run_font(run, "Calibri", 10.5 if level == 1 else (10 if level == 2 else 9.5),
                          bold=(level == 1 and entry_kind == "heading"), color=INK)
            run = paragraph.add_run("\t" + str(pages.get(text, "00")))
            _set_run_font(run, "Calibri", 10 if level == 1 else 9.5, color=INK)
        end_p = self.doc.add_paragraph()
        end = OxmlElement("w:fldChar")
        end.set(qn("w:fldCharType"), "end")
        run = end_p.add_run()
        run._r.append(end)
        end_p.paragraph_format.space_after = Pt(8)

    # -- couverture --------------------------------------------------------
    def cover(self, logo_path: str | Path, pages_total: str = ""):
        # En-tête institutionnel
        for text, size, bold, color in (
            ("RÉPUBLIQUE TUNISIENNE", 10.5, True, NAVY),
            ("Ministère de l'Enseignement Supérieur et de la Recherche Scientifique", 9, False, GREY),
            ("UNIVERSITÉ PRIVÉE TEK-UP", 11.5, True, NAVY),
            ("École Supérieure des Technologies de l'Information et de la Communication", 8.8, False, GREY),
        ):
            p = self.doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(1)
            run = p.add_run(text)
            _set_run_font(run, "Calibri", size, bold=bold, color=color)

        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run()
        run.add_picture(str(logo_path), width=Cm(4.2))

        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(18)
        run = p.add_run("Spécialité SSIR — Sécurité des Systèmes d'Information et Réseaux")
        _set_run_font(run, "Calibri", 9.2, color=GREY)

        # Titre entre deux filets
        title = self.doc.add_paragraph()
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        title.paragraph_format.space_before = Pt(6)
        title.paragraph_format.space_after = Pt(4)
        _border(title, top=True, color="10243F", size=12)
        run = title.add_run("CONCEPTION ET RÉALISATION\nD'UN SERVEUR MCP INTELLIGENT\nDE THREAT INTELLIGENCE")
        _set_run_font(run, "Calibri", 19, bold=True, color=NAVY)

        subtitle = self.doc.add_paragraph()
        subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
        subtitle.paragraph_format.space_after = Pt(6)
        run = subtitle.add_run("Protection du secret, étiquetage et traçabilité de bout en bout")
        _set_run_font(run, "Calibri", 12, italic=True, color=BLUE)
        _border(subtitle, bottom=True, color="10243F", size=12)

        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(22)
        run = p.add_run("Rapport de projet — « Art de protection du secret » & « Cyber OPS »")
        _set_run_font(run, "Calibri", 10.5, color=GREY)

        # Bloc d'identification
        info = self.doc.add_table(rows=4, cols=2)
        info.alignment = WD_TABLE_ALIGNMENT.CENTER
        info.autofit = False
        _cell_margins(info, 60, 60, 80, 80)
        labels = ["Réalisé par", "Encadrant académique", "Validation 1", "Validation 2"]
        values = [
            "Mzabi Sirine  ·  Khaled Ellif",
            "Tarek Hdji",
            "Architecture et exigences (P0 à P3)",
            "Réalisation, durcissement et tests (P4–P5)",
        ]
        for index, (label, value) in enumerate(zip(labels, values)):
            c0, c1 = info.rows[index].cells
            c0.width = Cm(5.0)
            c1.width = Cm(10.2)
            for cell, text, bold in ((c0, label, True), (c1, value, False)):
                paragraph = cell.paragraphs[0]
                paragraph.paragraph_format.space_after = Pt(0)
                run = paragraph.add_run(text)
                _set_run_font(run, "Calibri", 9.6, bold=bold, color=NAVY if bold else INK)

        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(28)
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run("Année universitaire 2026–2027")
        _set_run_font(run, "Calibri", 10.5, bold=True, color=NAVY)

        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run("Version 1.0 — 27 octobre 2026 · données synthétiques · laboratoire hors ligne")
        _set_run_font(run, "Calibri", 8, italic=True, color=GREY)

        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(16)
        run = p.add_run("Logo officiel : tek-up.de — document académique, usage non commercial.")
        _set_run_font(run, "Calibri", 7.2, italic=True, color=GREY)

    # -- sortie ------------------------------------------------------------
    def save(self, path: str | Path):
        if not self.emit:
            return
        Path(path).parent.mkdir(parents=True, exist_ok=True)
        self.doc.save(str(path))


def _split_code(text: str) -> list[str]:
    return text.replace("\r\n", "\n").split("\n")
